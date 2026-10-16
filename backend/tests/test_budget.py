"""Budget Upload (the old AbbottExecuteSPCreateAbtBudgetUpload) and Budget Report."""
import io

import openpyxl
import pytest

from tests.harness import Testing, client as c, ok, sign_in   # first: the test database
from app import models_master as X                             # noqa: E402
from app.routers.budget import COLUMNS                          # noqa: E402


@pytest.fixture(scope="module")
def h():
    return {"Authorization": f"Bearer {sign_in()}"}


@pytest.fixture(scope="module")
def world(h):
    db = Testing()
    try:
        country = db.query(X.Country).first()
        z = X.Zone(code="BZ", name="B ZONE", country_id=country.id, legacy_id=801)
        db.add(z); db.flush()
        db.add(X.Headquarter(code="BRT", name="BIRATNAGAR", zone_id=z.id, legacy_id=801))
        db.add(X.Headquarter(code="KTM", name="KATHMANDU", zone_id=z.id, legacy_id=802))
        db.add(X.ProductGroup(code="VN", name="VACCINE-NEURO", legacy_id=805))
        db.add(X.ProductMasterBrand(id=820, product_brand_name="EPILEX", product_group_id=805, reporting="X-OTHERS"))
        db.flush()
        db.add(X.ProductMaster(id=166, sap_product_name="EPILEX ORAL SOLUTION 200ML", product_brand_id=820, nrv=104.51))
        db.add(X.ProductMaster(id=167, sap_product_name="EPILEX CHRONO 300", product_brand_id=820, nrv=None))
        db.add(X.EmployeeMaster(id=29587, employee_ft_nm="RAJAN", employee_lt_nm="MANDAL"))
        db.add(X.YearMaster(year_number=2027))
        db.add(X.YearMaster(year_number=2028))
        db.add(X.EmployeeMaster(id=31796, employee_ft_nm="RAJU", employee_lt_nm="SHREEWASTAV"))
        db.add(X.ProductMaster(id=168, sap_product_name="EPILEX NEW 500", product_brand_id=820, nrv=50))
        db.add(X.EmployeeTerrAssignment(employee_id=31796, hq_id=801, territory_id=801))
        db.add(X.EmployeeTerrAssignment(employee_id=31796, hq_id=802, territory_id=802))
        db.commit()
    finally:
        db.close()
    return {}


def book(*rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row(emp=29587, group="VACCINE-NEURO", hq="BIRATNAGAR", brand="EPILEX", pid=166, name="EPILEX ORAL", units=(10,) * 12):
    return [emp, group, hq, brand, pid, name, *units]


def upload(h, data, year=2027, dry=True):
    return ok(c.post("/api/budget/upload", headers=h, data={"year": str(year), "dry_run": "true" if dry else "false"},
                     files=[("file", ("AbbottBudgetFormet.xlsx", data, "application/octet-stream"))]))


def test_the_check_names_every_bad_row_and_posts_nothing(h, world):
    r = upload(h, book(row(), row(emp=1), row(pid=999), row(pid=167), row(hq="NOWHERE"),
                       row(units=("x",) + (1,) * 11), row()), dry=False)
    assert r["message"].startswith("Nothing was posted")
    probs = {e["row"]: " ".join(e["problems"]) for e in r["errors"]}
    assert "EMPLOYEE_ID 1 is not in Employee" in probs[3]
    assert "PRODUCT_ID 999 is not in SKU" in probs[4]
    assert "has no NRV" in probs[5]
    assert "HQ_NAME NOWHERE is not in Head Quarter" in probs[6]
    assert "JANUARY_UNITS should be a number" in probs[7]
    assert "the same employee, HQ and SKU as row 2" in probs[8]
    assert ok(c.get("/api/budget/report?year=2027", headers=h))["rows"] == []


def test_a_good_file_is_posted_as_the_old_procedure_made_it(h, world):
    data = book(row(), row(hq="KATHMANDU", brand="OTHER BRAND", units=(5,) * 12))
    r = upload(h, data)
    assert (r["error_rows"], r["added"], r["updated"]) == (0, 2, 0)
    assert any("brand OTHER BRAND" in w for w in r["warnings"])
    assert ok(c.get("/api/budget/report?year=2027", headers=h))["rows"] == []        # checked only
    r = upload(h, data, dry=False)
    assert r["message"].startswith("Posted: 2 row(s) added")
    db = Testing()
    try:
        b = db.query(X.Budget).filter_by(target_year=2027, hq_name="BIRATNAGAR").one()
        assert (b.january_units, round(b.january, 2), b.product_group_id, b.product_brand_id, b.hq_id,
                b.product_brand_category, b.product_brand) == (10, 1045.10, 805, 820, 801, "X-OTHERS", "EPILEX")
        lines = {x.s_no for x in db.query(X.Budget).filter_by(target_year=2027)}
        assert len(lines) == 1                                    # one budget line: the SKU in the year
    finally:
        db.close()


def test_sending_again_updates_and_keeps_the_actuals(h, world):
    db = Testing()
    try:
        b = db.query(X.Budget).filter_by(target_year=2027, hq_name="BIRATNAGAR").one()
        b.january_actual = 0.55
        db.commit()
    finally:
        db.close()
    r = upload(h, book(row(units=(20,) * 12)), dry=False)
    assert (r["added"], r["updated"]) == (0, 1)
    db = Testing()
    try:
        b = db.query(X.Budget).filter_by(target_year=2027, hq_name="BIRATNAGAR").one()
        assert (b.january_units, round(b.january, 2), b.january_actual) == (20, 2090.2, 0.55)
    finally:
        db.close()


def test_the_report_is_the_old_one_in_inr(h, world):
    from app.routers.budget import REPORT_COLUMNS
    db = Testing()
    try:      # a row naming a SKU not in the master: the old report left such rows out
        db.add(X.Budget(s_no=99, target_year=2027, employee_id=29587, hq_id=801, product_id=99999))
        db.commit()
    finally:
        db.close()
    r = ok(c.get("/api/budget/report?year=2027", headers=h))
    assert (r["currency"], r["rate"], r["left_out"], len(r["rows"])) == ("INR", 1.6, 1, 2)
    ktm = [x for x in r["rows"] if x["hq_name"] == "KATHMANDU"][0]
    assert (ktm["employee_name"], ktm["hq_id"], ktm["product_group_id"], ktm["product_brand_id"]) == \
        ("RAJAN MANDAL", 802, 805, 820)
    assert ktm["units"][0] == 5 and round(ktm["values"][0], 4) == 326.5938      # 5 x 104.51 NPR / 1.6
    assert r["total_units"][0] == 25
    npr = ok(c.get("/api/budget/report?year=2027&hq_id=802&currency=NPR", headers=h))
    assert len(npr["rows"]) == 1 and round(npr["rows"][0]["values"][0], 2) == 522.55
    x = c.get("/api/budget/report/export?year=2027", headers=h)
    assert 'filename="BudgetReport_' in x.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert [c.value for c in ws[1]] == REPORT_COLUMNS
    assert ws.max_row == 4 and ws.cell(4, 1).value == "Total" and ws.cell(4, 12).value == 25
    assert c.get("/api/budget/report", headers=h).status_code == 422
    assert c.post("/api/budget/upload", headers=h, data={"year": "1999", "dry_run": "true"},
                  files=[("file", ("b.xlsx", book(row()), "x"))]).status_code == 422


def _zeros(year, **where):
    db = Testing()
    try:
        return db.query(X.Budget).filter_by(target_year=year, **where).all()
    finally:
        db.close()


def test_last_years_rows_missing_this_year_become_zero_rows(h, world):
    # 2027 has two budget rows (and one with a SKU not in the master); 2028 none
    r = ok(c.post("/api/budget/carry-forward", headers=h,
                  json={"from_year": 2027, "to_year": 2028, "dry_run": True}))
    assert (r["added"], r["already_there"], r["left_out"]) == (2, 0, {"SKU": 1})
    assert _zeros(2028) == []                                         # checked only
    r = ok(c.post("/api/budget/carry-forward", headers=h, json={"from_year": 2027, "to_year": 2028}))
    assert r["added"] == 2
    rows = _zeros(2028)
    assert {(x.hq_name, x.product_id) for x in rows} == {("BIRATNAGAR", 166), ("KATHMANDU", 166)}
    x = rows[0]
    assert (x.january_units, x.january, x.december_units, x.product_brand, x.product_group,
            x.product_brand_category) == (0, 0, 0, "EPILEX", "VACCINE-NEURO", "X-OTHERS")
    assert len({x.s_no for x in rows}) == 1                           # the SKU's one line in the year
    again = ok(c.post("/api/budget/carry-forward", headers=h, json={"from_year": 2027, "to_year": 2028}))
    assert (again["added"], again["already_there"]) == (0, 2)        # nothing twice


def test_carry_forward_for_one_employee_and_bad_years(h, world):
    r = ok(c.post("/api/budget/carry-forward", headers=h,
                  json={"from_year": 2027, "to_year": 2028, "employee_id": 31796, "dry_run": True}))
    assert r["added"] == 0                                            # had no 2027 budget
    assert c.post("/api/budget/carry-forward", headers=h,
                  json={"from_year": 2027, "to_year": 2027}).status_code == 422
    assert c.post("/api/budget/carry-forward", headers=h,
                  json={"from_year": 1999, "to_year": 2028}).status_code == 422


def test_a_new_product_gets_zero_rows_for_an_employee(h, world):
    # HQs left out: the HQs the employee covers on Employee Terr-PG (801 and 802)
    r = ok(c.post("/api/budget/zero-rows", headers=h,
                  json={"year": 2028, "employee_id": 31796, "product_ids": [168], "dry_run": True}))
    assert (r["added"], r["hq_ids"]) == (2, [801, 802])
    ok(c.post("/api/budget/zero-rows", headers=h, json={"year": 2028, "employee_id": 31796, "product_ids": [168]}))
    rows = _zeros(2028, employee_id=31796)
    assert sorted((x.hq_name, x.product_name, x.january_units) for x in rows) == \
        [("BIRATNAGAR", "EPILEX NEW 500", 0), ("KATHMANDU", "EPILEX NEW 500", 0)]
    # one HQ chosen, two SKUs, one already there
    r = ok(c.post("/api/budget/zero-rows", headers=h,
                  json={"year": 2028, "employee_id": 31796, "hq_ids": [801], "product_ids": [168, 166]}))
    assert (r["added"], r["already_there"]) == (1, 1)
    assert c.post("/api/budget/zero-rows", headers=h,
                  json={"year": 2028, "employee_id": 29587, "product_ids": [168]}).status_code == 422  # no HQ
    assert c.post("/api/budget/zero-rows", headers=h,
                  json={"year": 2028, "employee_id": 31796, "product_ids": []}).status_code == 422


def test_the_template_is_the_format(h):
    x = c.get("/api/budget/template", headers=h)
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert [c.value for c in ws[1]] == COLUMNS


def test_products_can_be_narrowed_by_group_and_brand(h, world):
    """A SKU with no group of its own is in its brand's group (EPILEX's are)."""
    db = Testing()
    try:
        gid = db.query(X.ProductGroup).filter_by(legacy_id=805).one().id
    finally:
        db.close()
    by_group = {o["id"] for o in ok(c.get(f"/api/masters/sku/options?product_group_id={gid}", headers=h))}
    by_brand = {o["id"] for o in ok(c.get("/api/masters/sku/options?product_brand_id=820", headers=h))}
    assert {166, 167, 168} <= by_group and by_brand == {166, 167, 168}
