"""Pharmachem: a zip of four Excel files, sheet 3 sales and sheet 4 returns.

The rule, as given:
 1. Distributor Name is always Pharmachem
 2. Invoice No read from the file
 3. Uploaded 1st-10th: Month is the previous month, else this month; in full
 4. Year is the current year, YYYY
 5. Uploaded 1st-11th: Mid Month N, else Y
 6-7. Customer and Product names from the file
 8. Quantity and Free Quantity from the file; positive on sheet 3, negative on sheet 4
 9. Rate from the file
11. B.Amount and Amount from the file; positive on sheet 3, negative on sheet 4
 -  Special characters omitted from product and customer names, but the '.' is
    kept in product names
"""
import io
import zipfile

import openpyxl
import pytest

from tests.harness import client as c, ok, sign_in
from app.services import parse
from app.services.rules import period
from datetime import date

HEAD = ["Invoice No", "Date", "Party Name", "Product Name", "Batch", "Qty", "Free",
        "Rate", "B.Amount", "Amount"]


def book(n: int, sheet_names=("Summary", "Stock", "Sales", "Sales Return")) -> bytes:
    """One of the four workbooks. Sheets 1-2 are things that must NOT be read."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_names[0]
    ws.append(["Party Name", "Product Name", "Qty", "Amount"])        # looks like sales: a trap
    ws.append(["TRAP PARTY", "TRAP PRODUCT", 999, 99999])
    st = wb.create_sheet(sheet_names[1])
    st.append(["Product Name", "Opening", "Closing"])
    st.append(["STOCK ONLY", 10, 20])

    s3 = wb.create_sheet(sheet_names[2])
    s3.append([f"PHARMACHEM PVT. LTD. - SALES REGISTER (FILE {n})"])
    s3.append([])
    s3.append(HEAD)
    s3.append([f"PC/{n}01", "05-10-2026", "Shree Medical Hall (P) Ltd.", "DUPHASTON 10MG TAB.",
               "B1", 10, 2, 120.5, 1205.0, 1265.25])
    # the same bill, second item: invoice and party written only on the first line
    s3.append([None, None, None, "UDILIV-300 TAB (10's)", "B2", 5, 0, 300, 1500, 1575])
    s3.append([f"PC/{n}02", "06-10-2026", "City Chemist & Co.", "DIGENE GEL MINT 200ML",
               "B3", 3, 1, 130.5, 391.5, 411.08])
    s3.append([None, None, None, "Total", None, 18, 3, None, 3096.5, 3251.33])  # left out
    s3.append([f"PC/{n}03", "07-10-2026", "Zero Line", "NOTHING SOLD", "B4", 0, 0, 10, 0, 0])

    s4 = wb.create_sheet(sheet_names[3])
    s4.append(HEAD)
    # written positive in the file; a return is negative in the sheet
    s4.append([f"SR/{n}01", "08-10-2026", "Shree Medical Hall (P) Ltd.", "DUPHASTON 10MG TAB.",
               "B1", 2, 0, 120.5, 241.0, 253.05])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def the_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n in range(1, 5):
            z.writestr(f"Pharmachem/PHARMA{n}.xlsx", book(n))
        z.writestr("Pharmachem/readme.txt", b"not a workbook")
    return buf.getvalue()


@pytest.fixture(scope="module")
def h():
    return {"Authorization": f"Bearer {sign_in()}"}


@pytest.fixture(scope="module")
def pharmachem(h):
    return [d for d in ok(c.get("/api/distributors", headers=h)) if d["name"] == "Pharmachem"][0]


# ------------------------------------------------------------- the rule
def test_the_seeded_rule(pharmachem):
    d = pharmachem
    assert d["cutoff_day"] == 10 and d["mid_cutoff_day"] == 11
    assert d["invoice_mode"] == "file" and d["rate_mode"] == "file"
    assert d["sales_sheet"] == 3 and d["return_sheet"] == 4 and d["product_keep_dot"]


@pytest.mark.parametrize("day,month,mid", [
    (1, "September", "N"), (10, "September", "N"),
    (11, "October", "N"),          # rule 3 says month, rule 5 still says N on the 11th
    (12, "October", "Y"), (28, "October", "Y")])
def test_month_and_mid_month_have_their_own_days(day, month, mid):
    p = period(date(2026, 10, day), 10, 11)
    assert (p["month"], p["year"], p["mid_month"]) == (month, 2026, mid)


def test_names_keep_the_dot_only_in_products():
    assert parse.clean("Shree Medical Hall (P) Ltd.") == "Shree Medical Hall P Ltd"
    assert parse.clean("DUPHASTON 10MG TAB.", keep=".") == "DUPHASTON 10MG TAB."
    assert parse.clean("UDILIV-300 TAB (10's)", keep=".") == "UDILIV 300 TAB 10 s"


# ------------------------------------------------------------ the reading
def test_only_sheet_3_and_sheet_4_are_read():
    r = parse.read_book("PHARMA1.xlsx", book(1), sales_sheet=3, return_sheet=4, keep_dot=True)
    assert r["error"] is None
    prods = {x["product"] for x in r["rows"]}
    assert "TRAP PRODUCT" not in prods and "STOCK ONLY" not in prods
    sales = [x for x in r["rows"] if not x["is_return"]]
    rets = [x for x in r["rows"] if x["is_return"]]
    assert len(sales) == 3 and len(rets) == 1
    assert r["omitted"] == 1                                  # the 0 + 0 line


def test_a_sheet_named_sheet3_is_taken_first():
    """Sheet3 named as such is read even when it is not third in order."""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet3"
    wb.active.append(HEAD)
    wb.active.append(["X1", None, "A", "PROD.A", None, 1, 0, 5, 5, 5])
    for t in ("Sheet1", "Sheet2", "Sheet4"):
        wb.create_sheet(t).append(HEAD)
    wb["Sheet1"].append(["WRONG", None, "B", "PROD B", None, 9, 0, 5, 45, 45])
    buf = io.BytesIO()
    wb.save(buf)
    r = parse.read_book("x.xlsx", buf.getvalue(), sales_sheet=3, return_sheet=4, keep_dot=True)
    assert [x["invoice"] for x in r["rows"]] == ["X1"]


def test_a_workbook_without_a_fourth_sheet_still_gives_its_sales():
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.create_sheet("B")
    s = wb.create_sheet("C")
    s.append(HEAD)
    s.append(["I1", None, "P", "X", None, 1, 0, 1, 1, 1])
    buf = io.BytesIO()
    wb.save(buf)
    r = parse.read_book("x.xlsx", buf.getvalue(), sales_sheet=3, return_sheet=4)
    assert len(r["rows"]) == 1 and r["error"] is None
    assert any("sheet 4" in s["name"] for s in r["sheets"])


# -------------------------------------------------------------- the journey
@pytest.fixture(scope="module")
def batch(h, pharmachem):
    r = ok(c.post("/api/uploads/inspect", headers=h, data={"distributor_id": pharmachem["id"]},
                  files=[("files", ("Pharmachem.zip", the_zip(), "application/zip"))]))
    ticked = sorted(f["name"] for f in r["files"] if f["suggested"])
    assert ticked == ["PHARMA1.xlsx", "PHARMA2.xlsx", "PHARMA3.xlsx", "PHARMA4.xlsx"]
    done = ok(c.post(f"/api/uploads/{r['token']}/commit", headers=h, json={
        "distributor_id": pharmachem["id"], "as_of": "2026-10-09", "files": ticked,
        "source": r["source"]}))
    rows = ok(c.get(f"/api/batches/{done['batch_id']}/rows?size=500", headers=h))["rows"]
    return done, rows


def test_four_files_become_one_sheet(batch):
    done, rows = batch
    assert done["rows"] == 16 == len(rows)              # 4 files x (3 sales + 1 return)
    assert {r["distributor_name"] for r in rows} == {"Pharmachem"}
    assert (done["month"], done["year"], done["mid_month"]) == ("September", 2026, "N")


def test_every_rule_holds_on_the_rows(batch):
    _, rows = batch
    first = [r for r in rows if r["invoice_no"] == "PC/101"]
    assert [r["product_name"] for r in first] == ["DUPHASTON 10MG TAB.", "UDILIV 300 TAB 10 s"]
    assert {r["customer_name"] for r in first} == {"Shree Medical Hall P Ltd"}   # filled down
    d = first[0]
    assert (d["quantity"], d["free_quantity"], d["rate"], d["b_amount"], d["amount"]) \
        == (10, 2, 120.5, 1205.0, 1265.25)
    ret = [r for r in rows if r["invoice_no"] == "SR/101"][0]
    assert ret["is_return"]
    assert (ret["quantity"], ret["free_quantity"], ret["b_amount"], ret["amount"]) \
        == (-2, 0, -241.0, -253.05)
    assert ret["rate"] == 120.5                         # a price stays positive
    sales = [r for r in rows if not r["is_return"]]
    assert all(r["quantity"] > 0 and r["amount"] > 0 for r in sales)
    assert not any(r["product_name"] in ("TRAP PRODUCT", "STOCK ONLY", "Total") for r in rows)


def test_the_download_has_the_twelve_columns(h, batch):
    done, _ = batch
    res = c.get(f"/api/batches/{done['batch_id']}/export", headers=h)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Sales"]
    assert [x.value for x in ws[1]] == ["Distributor Name", "Invoice No", "Month", "Year",
                                        "Mid Month", "Customer Name", "Product Name", "Quantity",
                                        "Free Quantity", "Rate", "B.Amount", "Amount"]
    assert ws.max_row == 17
    row2 = [x.value for x in ws[2]]
    assert row2[:5] == ["Pharmachem", "PC/101", "September", 2026, "N"]
    notes = {r[0].value: r[1].value for r in wb["Notes"].iter_rows() if r[0].value}
    assert "Sheet 3" in notes["Sheets read"] and "sheet 4" in notes["Sheets read"]


def test_the_11th_is_this_month_but_still_not_mid_month(h, pharmachem):
    r = ok(c.post("/api/uploads/inspect", headers=h, data={"distributor_id": pharmachem["id"]},
                  files=[("files", ("p.zip", the_zip(), "application/zip"))]))
    done = ok(c.post(f"/api/uploads/{r['token']}/commit", headers=h, json={
        "distributor_id": pharmachem["id"], "as_of": "2026-10-11",
        "files": [f["name"] for f in r["files"] if f["suggested"]]}))
    assert (done["month"], done["mid_month"]) == ("October", "N")
