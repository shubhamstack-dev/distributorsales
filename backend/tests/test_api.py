"""Rules, parsing, the wall, and the numbers coming back out the same."""
import io
import re
from datetime import date

import openpyxl
import pytest
from sqlalchemy import select

from tests.harness import client as c, anon, Testing, ok, sign_in
from app import models as M
from app.services.rules import period
from app.services.parse import clean, read_book, books_in

ZIP = "/mnt/user-data/uploads/FW__MID-MONTH_SALES.zip"
PW = ["AHCPW.xlsx", "AILGPW.xlsx", "AILNPW.xlsx"]


@pytest.fixture(scope="module")
def tok():
    return sign_in()


def H(t):
    return {"Authorization": f"Bearer {t}"}


def dist(tok, name="Yetichem"):
    return [d for d in ok(c.get("/api/distributors", headers=H(tok))) if d["name"] == name][0]


# ================================================================ the wall
def test_nothing_answers_without_a_token():
    open_ = []
    for path, ops in c.app.openapi()["paths"].items():
        for method in ops:
            if (method.upper(), path) in {("GET", "/api/health"), ("POST", "/api/login")}:
                continue
            r = anon.request(method.upper(), re.sub(r"\{[^}]+\}", "1", path), json={})
            if r.status_code != 401:
                open_.append((method.upper(), path, r.status_code))
    assert open_ == [], open_


def test_a_wrong_password_is_refused():
    assert c.post("/api/login", json={"password": "nope"}).status_code == 401


def test_a_tampered_token_is_refused(tok):
    body, sig = tok.split(".")
    assert c.get("/api/batches", headers=H(f"{body}.{sig[:-2]}xx")).status_code == 401


# ================================================================ the rules
@pytest.mark.parametrize("day,cut,month,year,mid", [
    ("2026-09-25", 15, "September", 2026, "Y"),
    ("2026-09-15", 15, "August", 2026, "N"),
    ("2026-09-01", 15, "August", 2026, "N"),
    ("2026-09-16", 15, "September", 2026, "Y"),
    ("2026-04-10", 10, "March", 2026, "N"),     # Pharmachem's cut-off
    ("2026-04-11", 10, "April", 2026, "Y"),
    ("2026-12-20", 15, "December", 2026, "Y"),
])
def test_the_month_rule(day, cut, month, year, mid):
    p = period(date.fromisoformat(day), cut)
    assert (p["month"], p["year"], p["mid_month"]) == (month, year, mid)


def test_january_keeps_the_current_year_and_says_so():
    """The rule says the year is the current year. It is followed literally, and
    the odd case is flagged rather than quietly corrected."""
    p = period(date(2026, 1, 10), 15)
    assert p["month"] == "December" and p["year"] == 2026 and p["wrapped"] is True


# =============================================================== the parser
def test_special_characters_become_spaces():
    assert clean("Rantoks Pharma Pvt.ltd.") == "Rantoks Pharma Pvt ltd"
    assert clean("Mahendra Medical Pvt. Ltd") == "Mahendra Medical Pvt Ltd"
    assert clean("STEMETIL 5MG TAB ") == "STEMETIL 5MG TAB"


def test_the_zip_is_opened_and_only_workbooks_taken():
    data = open(ZIP, "rb").read()
    books = books_in("FW__MID-MONTH_SALES.zip", data)
    assert len(books) == 10
    assert all(n.endswith(".xlsx") for n, _ in books)


def test_the_cross_tab_reads_to_the_same_numbers():
    data = open(ZIP, "rb").read()
    rows, omitted, amount = 0, 0, 0.0
    for name, blob in books_in("z.zip", data):
        if name not in PW:
            continue
        r = read_book(name, blob)
        rows += len(r["rows"])
        omitted += r["omitted"]
        amount += sum(x["amount"] for x in r["rows"])
        assert r["sheets"][0]["layout"] == "cross-tab"
    assert (rows, omitted) == (82, 140)
    assert round(amount, 2) == 2765117.77


def test_a_stocks_file_is_refused_rather_than_half_read():
    data = open(ZIP, "rb").read()
    blob = dict(books_in("z.zip", data))["AHCSTOCKS.xlsx"]
    r = read_book("AHCSTOCKS.xlsx", blob)
    assert r["rows"] == [] and "No sales rows" in r["error"]


# ============================================================== the journey
@pytest.fixture(scope="module")
def batch(tok):
    d = dist(tok)
    with open(ZIP, "rb") as fh:
        r = ok(c.post("/api/uploads/inspect", headers=H(tok),
                      data={"distributor_id": d["id"]},
                      files=[("files", ("FW__MID-MONTH_SALES.zip", fh.read(), "application/zip"))]))
    assert len(r["files"]) == 10
    suggested = [f["name"] for f in r["files"] if f["suggested"]]
    assert sorted(suggested) == PW              # the rule picks its own files
    done = ok(c.post(f"/api/uploads/{r['token']}/commit", headers=H(tok), json={
        "distributor_id": d["id"], "as_of": "2026-09-25", "files": suggested,
        "source": r["source"], "who": "tester"}))
    return done


def test_the_batch_matches_the_earlier_numbers(batch):
    assert batch["rows"] == 82 and batch["omitted"] == 140
    assert batch["total_qty"] == 43200 and batch["total_amount"] == 2765117.77
    assert (batch["month"], batch["year"], batch["mid_month"]) == ("September", 2026, "Y")


def test_the_batch_records_what_was_left_out(tok, batch):
    b = ok(c.get(f"/api/batches/{batch['batch_id']}", headers=H(tok)))
    used = [f["name"] for f in b["files"] if f["used"]]
    left = [f["name"] for f in b["files"] if not f["used"]]
    assert sorted(used) == PW and len(left) == 7
    assert "AHCSTOCKS.xlsx" in left


def test_every_rule_holds_on_the_stored_rows(tok, batch):
    got = ok(c.get(f"/api/batches/{batch['batch_id']}/rows?size=500", headers=H(tok)))
    rows = got["rows"]
    assert got["total"] == 82
    assert all(r["distributor_name"] == "Yetichem" for r in rows)
    assert [r["invoice_no"] for r in rows[:3]] == ["INV - 1", "INV - 2", "INV - 3"]
    assert rows[-1]["invoice_no"] == "INV - 82"
    assert all(r["month"] == "September" and r["year"] == 2026 and r["mid_month"] == "Y" for r in rows)
    assert not any(re.search(r"[^A-Za-z0-9 ]", r["customer_name"] + r["product_name"]) for r in rows)
    assert not any(r["quantity"] == 0 and r["free_quantity"] == 0 for r in rows)
    assert all(abs(r["rate"] * r["quantity"] - r["amount"]) < 0.005 for r in rows if r["quantity"])
    assert all(r["b_amount"] == r["amount"] for r in rows)


def test_the_rows_can_be_searched(tok, batch):
    got = ok(c.get(f"/api/batches/{batch['batch_id']}/rows?q=STEMETIL", headers=H(tok)))
    assert got["total"] > 0
    assert all("STEMETIL" in r["product_name"] for r in got["rows"])


def test_the_export_is_the_sheet_that_was_asked_for(tok, batch):
    r = c.get(f"/api/batches/{batch['batch_id']}/export", headers=H(tok))
    assert r.status_code == 200
    assert "Yetichem_Sales_September2026.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    assert wb.sheetnames == ["Sales", "Notes"]
    ws = wb["Sales"]
    assert [c.value for c in ws[1]] == [
        "Distributor Name", "Invoice No", "Month", "Year", "Mid Month", "Customer Name",
        "Product Name", "Quantity", "Free Quantity", "Rate", "B.Amount", "Amount"]
    assert ws.max_row == 83
    assert round(sum(ws.cell(r_, 12).value for r_ in range(2, 84)), 2) == 2765117.77
    notes = {wb["Notes"].cell(i, 1).value: wb["Notes"].cell(i, 2).value
             for i in range(1, wb["Notes"].max_row + 1)}
    assert "AHCPW.xlsx" in notes["Files converted"]
    assert "AHCSTOCKS.xlsx" in notes["Files offered but not converted"]


def test_choosing_nothing_is_refused(tok):
    d = dist(tok)
    with open(ZIP, "rb") as fh:
        r = ok(c.post("/api/uploads/inspect", headers=H(tok), data={"distributor_id": d["id"]},
                      files=[("files", ("z.zip", fh.read(), "application/zip"))]))
    bad = c.post(f"/api/uploads/{r['token']}/commit", headers=H(tok),
                 json={"distributor_id": d["id"], "as_of": "2026-09-25", "files": []})
    assert bad.status_code == 422


def test_choosing_only_unreadable_files_is_refused(tok):
    d = dist(tok)
    with open(ZIP, "rb") as fh:
        r = ok(c.post("/api/uploads/inspect", headers=H(tok), data={"distributor_id": d["id"]},
                      files=[("files", ("z.zip", fh.read(), "application/zip"))]))
    bad = c.post(f"/api/uploads/{r['token']}/commit", headers=H(tok), json={
        "distributor_id": d["id"], "as_of": "2026-09-25", "files": ["AHCSTOCKS.xlsx"]})
    assert bad.status_code == 409 and "Nothing could be read" in bad.text


def test_an_expired_upload_says_so(tok):
    d = dist(tok)
    bad = c.post("/api/uploads/deadbeef/commit", headers=H(tok),
                 json={"distributor_id": d["id"], "files": ["x.xlsx"]})
    assert bad.status_code == 410


# ============================================================ distributors
def test_the_three_are_seeded_with_their_own_rules(tok):
    ds = {d["name"]: d for d in ok(c.get("/api/distributors", headers=H(tok)))}
    assert ds["Yetichem"]["cutoff_day"] == 15 and ds["Yetichem"]["invoice_mode"] == "seq"
    assert ds["Pharmachem"]["cutoff_day"] == 10 and ds["Pharmachem"]["invoice_mode"] == "file"
    assert ds["Gunjeshwari"]["cutoff_day"] == 10 and ds["Gunjeshwari"]["invoice_mode"] == "seq"


def test_a_rule_can_be_corrected_without_a_deployment(tok):
    d = dist(tok, "Gunjeshwari")
    out = ok(c.put(f"/api/distributors/{d['id']}", headers=H(tok), json={"cutoff_day": 12}))
    assert out["cutoff_day"] == 12
    ok(c.put(f"/api/distributors/{d['id']}", headers=H(tok), json={"cutoff_day": 10}))


def test_a_nonsense_rule_is_refused(tok):
    d = dist(tok, "Gunjeshwari")
    assert c.put(f"/api/distributors/{d['id']}", headers=H(tok),
                 json={"cutoff_day": 31}).status_code == 422
    assert c.put(f"/api/distributors/{d['id']}", headers=H(tok),
                 json={"pick_pattern": "([unclosed"}).status_code == 422


def test_a_batch_can_be_deleted_and_takes_its_rows(tok):
    d = dist(tok)
    with open(ZIP, "rb") as fh:
        r = ok(c.post("/api/uploads/inspect", headers=H(tok), data={"distributor_id": d["id"]},
                      files=[("files", ("z.zip", fh.read(), "application/zip"))]))
    made = ok(c.post(f"/api/uploads/{r['token']}/commit", headers=H(tok), json={
        "distributor_id": d["id"], "as_of": "2026-09-25",
        "files": [f["name"] for f in r["files"] if f["suggested"]]}))
    bid = made["batch_id"]
    ok(c.delete(f"/api/batches/{bid}", headers=H(tok)), 204)
    assert c.get(f"/api/batches/{bid}", headers=H(tok)).status_code == 404
    db = Testing()
    left = db.execute(select(M.SalesRow).where(M.SalesRow.batch_id == bid)).first()
    db.close()
    assert left is None
