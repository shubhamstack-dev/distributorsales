"""Gunjeshwari end to end: two batch-wise PDFs combined into the twelve-column
sheet, Rate read from the PDF, names replaced from the reference lists."""
import io
import zipfile

import openpyxl
import pytest
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from tests.harness import client as c, Testing, ok, sign_in
from app import models as M
from app.services import names, seed
from app.services.excel import HEADS

HEAD = ["Party Name", "Product Name", "Batch", "Qty", "Free", "Rate", "Amount"]


def _pdf(rows):
    buf = io.BytesIO()
    t = Table([HEAD] + rows, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(buf, pagesize=A4).build([t])
    return buf.getvalue()


# ALI: one party; the second product runs over two batches, the second batch
# line starting from the quantity with the product left blank.
ALI = _pdf([
    ["Shree Medical Hall (P) Ltd.", "", "", "", "", "", ""],          # party heading
    ["", "Paracetamol-500 Tab.", "PB01", "10", "2", "12.50", "125.00"],
    ["", "Amoxy 250 Cap", "AX01", "5", "0", "40.00", "200.00"],
    ["", "", "AX02", "3", "1", "40.00", "120.00"],                    # same product, new batch
    ["", "Total", "", "18", "3", "", "445.00"],
])
# AHL: two parties written on every line, one name not on the customer list.
AHL = _pdf([
    ["OM PHARMACY", "ORS Sachet", "OR9", "20", "0", "8.25", "165.00"],
    ["New Hope Clinic", "ORS Sachet", "OR9", "4", "0", "8.25", "33.00"],
    ["Grand Total", "", "", "24", "0", "", "198.00"],
])


def _zip():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Gunjeshwari/ALI BATCHWISE.pdf", ALI)
        z.writestr("Gunjeshwari/AHL BATCHWISE.PDF", AHL)
        z.writestr("Gunjeshwari/Outstanding.pdf", AHL)              # not a batch report
    return buf.getvalue()


def _xlsx(head, values):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sr", head])
    for i, v in enumerate(values, 1):
        ws.append([i, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CUSTOMERS = _xlsx("Customer Name", ["SHREE MEDICAL HALL P LTD", "Om Pharmacy", "Om Pharmacy."])
PRODUCTS = _xlsx("Product Name", ["PARACETAMOL 500 TAB", "AMOXY 250MG CAP", "O.R.S. SACHET"])


@pytest.fixture(scope="module")
def tok():
    return sign_in()


def H(t):
    return {"Authorization": f"Bearer {t}"}


@pytest.fixture(scope="module")
def gunj(tok):
    return [d for d in ok(c.get("/api/distributors", headers=H(tok)))
            if d["name"] == "Gunjeshwari"][0]


# ================================================================ name lists
def test_names_match_ignoring_case_spaces_and_special_characters():
    assert names.norm("Shree Medical Hall (P) Ltd.") == names.norm("SHREE MEDICAL HALL P LTD")
    assert names.norm("Paracetamol-500 Tab.") == names.norm("PARACETAMOL 500 TAB")
    m = names.Mapper({names.norm("Om Pharmacy"): "Om Pharmacy"})
    assert m("OM PHARMACY") == "Om Pharmacy" and m("New Hope Clinic") == "New Hope Clinic"
    assert m.missed == {"New Hope Clinic": 1} and "New Hope Clinic" in m.summary("customer")


def test_a_list_is_read_from_its_header_or_its_first_text_column():
    assert names.read_names(CUSTOMERS, "customer") == [
        "SHREE MEDICAL HALL P LTD", "Om Pharmacy", "Om Pharmacy."]
    wb = openpyxl.Workbook()
    for v in ["Alpha Stores", "Beta Medicos"]:
        wb.active.append([v])
    buf = io.BytesIO()
    wb.save(buf)
    assert names.read_names(buf.getvalue(), "customer") == ["Alpha Stores", "Beta Medicos"]


def test_uploading_a_list_replaces_the_old_one(tok, gunj):
    url = f"/api/distributors/{gunj['id']}/names"
    first = ok(c.post(url, headers=H(tok), data={"kind": "customer"},
                      files={"file": ("old.xlsx", _xlsx("Customer Name", ["Gone Ltd"]), "x")}))
    assert first["customer"]["count"] == 1
    r = ok(c.post(url, headers=H(tok), data={"kind": "customer"},
                  files={"file": ("customers.xlsx", CUSTOMERS, "x")}))
    # 'Om Pharmacy' and 'Om Pharmacy.' are the same name once matched
    assert r["read"] == 3 and r["kept"] == 2 and r["customer"]["file"] == "customers.xlsx"
    r = ok(c.post(url, headers=H(tok), data={"kind": "product"},
                  files={"file": ("products.xlsx", PRODUCTS, "x")}))
    assert r["product"]["count"] == 3
    assert c.post(url, headers=H(tok), data={"kind": "customer"},
                  files={"file": ("bad.xlsx", b"nope", "x")}).status_code == 422


# ============================================================== end to end
@pytest.fixture(scope="module")
def batch(tok, gunj):
    test_uploading_a_list_replaces_the_old_one(tok, gunj)
    r = ok(c.post("/api/uploads/inspect", headers=H(tok),
                  files=[("files", ("Gunjeshwari.zip", _zip(), "application/zip"))],
                  data={"distributor_id": str(gunj["id"])}))
    picked = sorted(f["name"] for f in r["files"] if f["suggested"])
    assert picked == ["AHL BATCHWISE.PDF", "ALI BATCHWISE.pdf"]
    b = ok(c.post(f"/api/uploads/{r['token']}/commit", headers=H(tok), json={
        "distributor_id": gunj["id"], "as_of": "2026-10-05", "files": picked,
        "source": r["source"]}))
    rows = ok(c.get(f"/api/batches/{b['batch_id']}/rows?size=50", headers=H(tok)))["rows"]
    return b, rows


def test_both_batchwise_pdfs_are_combined_into_one_sequence(batch):
    b, rows = batch
    assert b["rows"] == 5 and len(rows) == 5          # 3 from ALI, 2 from AHL
    assert [r["invoice_no"] for r in rows] == [f"INV - {i}" for i in range(1, 6)]
    assert {r["source_file"] for r in rows} == {"AHL BATCHWISE.PDF", "ALI BATCHWISE.pdf"}
    assert all(r["distributor_name"] == "Gunjeshwari" for r in rows)
    assert b["total_qty"] == 42 and b["total_amount"] == 643.00


def test_the_month_rule_and_year(batch):
    b, rows = batch
    # uploaded on the 5th: on or before the 10th, so the month before and N
    assert all((r["month"], r["year"], r["mid_month"]) == ("September", 2026, "N") for r in rows)


def test_a_product_repeated_from_its_quantity_is_its_own_row(batch):
    _, rows = batch
    amoxy = [r for r in rows if r["product_name"] == "Amoxy 250 Cap"]
    assert [(r["quantity"], r["free_quantity"], r["amount"]) for r in amoxy] == [
        (5.0, 0.0, 200.0), (3.0, 1.0, 120.0)]
    assert all(r["customer_name"] == "SHREE MEDICAL HALL P LTD" for r in amoxy)


def test_rate_is_read_from_the_file_and_b_amount_equals_amount(batch):
    _, rows = batch
    ors = [r for r in rows if r["product_name"] == "O.R.S. SACHET"]
    assert [r["rate"] for r in ors] == [8.25, 8.25]
    assert all(r["b_amount"] == r["amount"] for r in rows)


def test_names_are_replaced_from_the_lists_and_misses_reported(batch):
    b, rows = batch
    assert {r["product_name"] for r in rows} == {"PARACETAMOL 500 TAB", "Amoxy 250 Cap",
                                                 "O.R.S. SACHET"}
    # 'AMOXY 250 CAP' is not 'AMOXY 250MG CAP', so it stays as extracted
    assert b["unmatched_products"] == ["Amoxy 250 Cap"]
    assert b["unmatched_customers"] == ["New Hope Clinic"]
    assert {r["customer_name"] for r in rows} == {"SHREE MEDICAL HALL P LTD", "Om Pharmacy",
                                                  "New Hope Clinic"}


def test_the_export_has_the_twelve_headings_and_the_notes(tok, batch):
    b, _ = batch
    res = c.get(f"/api/batches/{b['batch_id']}/export", headers=H(tok))
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Sales"]
    assert [x.value for x in ws[1]] == HEADS == [
        "Distributor Name", "Invoice No", "Month", "Year", "Mid Month", "Customer Name",
        "Product Name", "Quantity", "Free Quantity", "Rate", "B.Amount", "Amount"]
    assert [x.value for x in ws[2]][:2] == ["Gunjeshwari", "INV - 1"]
    notes = {r[0].value: r[1].value for r in wb["Notes"].iter_rows()}
    assert notes["Rate"].startswith("Read from the file")
    assert "New Hope Clinic" in notes["Name mapping and notes"]


def test_after_the_tenth_it_is_the_current_month_and_y(tok, gunj):
    r = ok(c.post("/api/uploads/inspect", headers=H(tok),
                  files=[("files", ("g.zip", _zip(), "application/zip"))],
                  data={"distributor_id": str(gunj["id"])}))
    b = ok(c.post(f"/api/uploads/{r['token']}/commit", headers=H(tok), json={
        "distributor_id": gunj["id"], "as_of": "2026-10-11",
        "files": ["ALI BATCHWISE.pdf", "AHL BATCHWISE.PDF"]}))
    assert (b["month"], b["year"], b["mid_month"]) == ("October", 2026, "Y")


def test_an_old_database_gets_rate_from_file_with_the_new_note():
    db = Testing()
    try:
        d = db.query(M.Distributor).filter_by(name="Gunjeshwari").one()
        keep = (d.note, d.rate_mode)
        d.note, d.rate_mode = seed._GUNJESHWARI_OLD_NOTES[1], "calc"
        db.commit()
        seed.update_rules(db)
        assert d.note == seed.GUNJESHWARI_NOTE and d.rate_mode == "file"
        d.note, d.rate_mode = "typed by hand", "calc"
        db.commit()
        seed.update_rules(db)
        assert (d.note, d.rate_mode) == ("typed by hand", "calc")
        d.note, d.rate_mode = keep
        db.commit()
    finally:
        db.close()
