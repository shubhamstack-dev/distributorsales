"""The old system's data tables: Sales (MAPPING_MASTER), Mid sales
(MAPPING_MASTER_MID), Budget (BUDGET_MASTER) and the four *_ASSIGNMENT tables,
uploaded from one workbook the way Book1.xlsx has them."""
import io

import openpyxl
import pytest

from tests.harness import client as c, ok, sign_in   # first: it points the app at a test database
from app import models_master as X                     # noqa: E402


@pytest.fixture(scope="module")
def h():
    return {"Authorization": f"Bearer {sign_in()}"}


def wb_bytes(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SALES_H = [n for n, _ in X.SALES_COLUMNS]
BUDGET_H = [n for n, _ in X.BUDGET_COLUMNS if n != "ID"]


def sale(i, cust="ENERGY PHARMA SALES", qty=200, month="January", invoice=83038554):
    v = {"ID": i, "DIST_ID": 12664157, "PRODUCT_NAME": "INDERAL TABS 20MG", "PRODUCT_ID": 154,
         "PRODUCT_GROUP_ID": 5, "RATE": 45.12, "HQ_ID": 6, "MONTH": month, "YEAR": 2016,
         "TERRITORY_ID": 6, "PRODUCT_BRAND_ID": 18, "DIST_NAME": "GUNJESHWARI", "STATUS": "NULL",
         "SALES_QTY": qty, "CUSTOMER_NAME": cust, "CUSTOMER_MASTER_NAME": cust,
         "PRODUCT_GROUP_NAME": "VACCINE-NEURO", "HQ_NAME": "KATHMANDU", "TERRITORY_NAME": "KATHMANDU",
         "PRODUCT_BRAND_NAME": "INDERAL", "INVOICE_NO": invoice, "CUSTOMER_ID": 45, "FREE_QUANTITY": 0,
         "RETURN_QUANTITY": 0, "PRODUCT_MASTER_NAME": "INDERAL TABS 20MG",
         "PRODUCT_BRAND_CATEGORY": "X-OTHERS"}
    return [v.get(n, "NULL") for n in SALES_H]


def budget(s_no, emp, hq_id, hq, jan_actual=1.25415):
    v = {"S_NO": s_no, "EMPLOYEE_ID": emp, "PRODUCT_GROUP": "VACCINE-NEURO", "PRODUCT_BRAND": "EPILEX",
         "HQ_NAME": hq, "PRODUCT_ID": 166, "PRODUCT_NAME": "EPILEX ORAL SOLUTION 200ML",
         "TARGET_YEAR": 2023, "JANUARY_ACTUAL": jan_actual, "PRODUCT_GROUP_ID": 5,
         "PRODUCT_BRAND_ID": 20, "HQ_ID": hq_id, "PRODUCT_BRAND_CATEGORY": "X-OTHERS"}
    return [v.get(n, 0) for n in BUDGET_H]


def book(**over) -> bytes:
    sheets = {
        "BUDGET_MASTER_TEST_NEW": [BUDGET_H, budget(40293, 29587, 1, "BIRATNAGAR"),
                                   # S_NO repeats across HQs: a different row each time
                                   budget(40293, 29587, 3, "BUTWAL", jan_actual="NULL")],
        "MAPPING_MASTER": [SALES_H, sale(324779), sale(324780, cust="ARCHANA MEDICINE CENTRE ",
                                                       invoice="ABP77-00629")],
        "MAPPING_MASTER_MID": [SALES_H + ["MID"], sale(1, month="FEBRUARY") + ["NULL"]],
        "CUSTOMER_TERR_ASSIGNMENT": [["CUSTOMER_ID", "TERRITORY_ID", "HQ_ID", "PRODUCT_GROUP_ID", "TEAM_ID"],
                                     [1, 3, 3, 1, 5], [1, 3, 3, 2, 1]],
        "CUSTOMER_PG_ASSIGNMENT": [["CUSTOMER_ID", "PRODUCT_GROUP_ID", "PRODUCT_BRAND_ID", "PRODUCT_ID"],
                                   [1, 1, 2, 31], [1, 1, 6, 49]],
        "EMPLOYEE_TERR_ASSIGNMENT": [["EMPLOYEE_ID", "HQ_ID", "TERRITORY_ID"], [29587, 1, 1]],
        "EMPLOYEE_PG_ASSIGNMENT": [["EMPLOYEE_ID", "PRODUCT_ID", "PRODUCT_BRAND_ID", "PRODUCT_GROUP_ID"],
                                   [29025, 4, 13, 8], [29025, 5, 13, 8]],
    }
    sheets.update(over)
    return wb_bytes(sheets)


def upload(h, data, dry=True, slug=None):
    form = {"dry_run": "true" if dry else "false"}
    if slug:
        form["slug"] = slug
    return ok(c.post("/api/masters/upload", headers=h, data=form,
                     files=[("files", ("Book1.xlsx", data, "application/octet-stream"))]))


def by_slug(r):
    return {t["slug"]: t for t in r["tables"]}


def test_every_sheet_of_the_workbook_finds_its_table(h):
    r = upload(h, book())
    t = by_slug(r)
    assert set(t) == {"budget", "sales", "mid_sales", "customer_terr_assignment",
                      "customer_pg_assignment", "employee_terr_assignment", "employee_pg_assignment"}
    assert {k: v["added"] for k, v in t.items()} == {
        "budget": 2, "sales": 2, "mid_sales": 1, "customer_terr_assignment": 2,
        "customer_pg_assignment": 2, "employee_terr_assignment": 1, "employee_pg_assignment": 2}
    assert not any(v["skipped"] for v in t.values())
    # a check writes nothing
    assert ok(c.get("/api/masters/sales", headers=h))["total"] == 0


def test_import_keeps_the_values_and_a_second_import_adds_nothing(h):
    upload(h, book(), dry=False)
    rows = ok(c.get("/api/masters/sales?size=10", headers=h))["rows"]
    s = {r["id"]: r for r in rows}
    assert s[324779]["rate"] == 45.12 and s[324779]["month"] == "January"
    assert s[324779]["invoice_no"] == "83038554" and s[324780]["invoice_no"] == "ABP77-00629"
    assert s[324779]["status"] is None                        # "NULL" in the sheet is empty
    mid = ok(c.get("/api/masters/mid_sales", headers=h))["rows"][0]
    assert mid["month"] == "February"                         # FEBRUARY read as the month's name
    b = ok(c.get("/api/masters/budget?size=10", headers=h))["rows"]
    assert sorted((x["s_no"], x["hq_id"], x["january_actual"]) for x in b) == \
        [(40293, 1, 1.25415), (40293, 3, None)]
    r = upload(h, book(), dry=False)
    assert all(t["added"] == 0 and t["updated"] == 0 and t["unchanged"] == t["rows"]
               for t in r["tables"]), r
    assert ok(c.get("/api/masters/customer_pg_assignment", headers=h))["total"] == 2


def test_a_changed_row_is_updated_in_place(h):
    data = book(MAPPING_MASTER=[SALES_H, sale(324779, qty=250)])
    t = by_slug(upload(h, data, dry=False))["sales"]
    assert (t["added"], t["updated"]) == (0, 1)
    rows = ok(c.get("/api/masters/sales?size=10", headers=h))["rows"]
    assert {r["id"]: r["sales_qty"] for r in rows} == {324779: 250, 324780: 200}


def test_a_bad_row_is_named_and_the_rest_go_in(h):
    data = book(MAPPING_MASTER=[SALES_H, sale(900001, qty="lots"), sale(900002, month="Smarch"),
                                sale(900003)])
    t = by_slug(upload(h, data))["sales"]
    assert (t["added"], t["skipped"]) == (1, 2)
    assert "SALES_QTY should be a number" in t["errors"][0]
    assert "MONTH has to be one of" in t["errors"][1]


def test_budget_needs_its_key_columns(h):
    head = [x for x in BUDGET_H if x != "HQ_ID"]
    row = [v for x, v in zip(BUDGET_H, budget(1, 2, 3, "X")) if x != "HQ_ID"]
    t = by_slug(upload(h, book(BUDGET_MASTER_TEST_NEW=[head, row])))["budget"]
    assert t["added"] == 0 and "no HQ_ID column" in t["errors"][0]


def test_ids_not_in_the_masters_are_kept_and_reported_once(h):
    t = by_slug(upload(h, book()))["customer_terr_assignment"]
    assert any(w.startswith("CUSTOMER_ID 1: not in Customer/Hospital") for w in t["warnings"]), t["warnings"]


def test_the_lists_filter_as_the_screens_ask(h):
    assert ok(c.get("/api/masters/sales?year=2016&month=January&dist_name=gunjeshwari",
                    headers=h))["total"] == 2
    assert ok(c.get("/api/masters/sales?dist_name=YETICHEM", headers=h))["total"] == 0
    assert ok(c.get("/api/masters/budget?target_year=2023&hq_id=3", headers=h))["total"] == 1
    assert ok(c.get("/api/masters/customer_terr_assignment?team_id=5", headers=h))["total"] == 1


def test_an_assignment_added_on_screen_cannot_be_added_twice(h):
    body = {"employee_id": 31467, "hq_id": 2, "territory_id": 2}
    ok(c.post("/api/masters/employee_terr_assignment", headers=h, json=body), 201)
    r = c.post("/api/masters/employee_terr_assignment", headers=h, json=body)
    assert r.status_code == 409 and "already there" in r.json()["detail"]


def test_export_writes_the_workbook_columns_and_reads_back_in(h):
    res = c.get("/api/masters/sales/export", headers=h)
    assert res.status_code == 200 and 'filename="MAPPING_MASTER.xlsx"' in res.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(res.content), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == SALES_H and len(rows) == 3
    r = upload(h, res.content, slug="sales")
    assert r["tables"][0]["unchanged"] == 2
    # the data tables are not masters, so "every master" leaves them out
    wb = openpyxl.load_workbook(io.BytesIO(c.get("/api/masters/export-all", headers=h).content),
                                read_only=True)
    assert "MAPPING_MASTER" not in wb.sheetnames and "CUSTOMER_MASTER" in wb.sheetnames


def test_the_screens_are_listed_under_assignment_and_sales(h):
    meta = {m["slug"]: m for m in ok(c.get("/api/masters/meta", headers=h))}
    # the four tables sit behind the joined Terr-PG pages, not in the menu themselves
    assert [m["title"] for m in meta.values() if m["group"] == "Assignment"] == ["Team PG HQ"]
    assert [m["slug"] for m in meta.values() if m["group"] == "Assignment tables"] == [
        "customer_terr_assignment", "customer_pg_assignment", "employee_terr_assignment",
        "employee_pg_assignment"]
    assert not any(m["group"] == "Alignment" for m in meta.values())
    assert [meta[s]["title"] for s in ("sales", "mid_sales", "budget")] == ["Sales", "Mid sales", "Budget"]
    assert meta["sales"]["table"] == "MAPPING_MASTER" and meta["sales"]["filters"][0] == "year"


# ------------------------------------------- Customer / Employee Terr-PG pages
def test_customer_terr_pg_joins_on_customer_and_product_group(h):
    r = ok(c.get("/api/terr-pg/customer", headers=h))
    assert r["columns"] == ["Customer Name", "HQ", "Territory", "Product Group", "Team",
                            "Product Brand", "Product"]
    # both product rows of customer 1 are in group 1, covered by territory 3 / team 5;
    # the group-2 territory row has no products and does not appear
    assert r["total"] == 2
    assert [(x[0], x[1], x[2], x[3], x[4], x[6]) for x in r["rows"]] == \
        [(1, 3, 3, 1, 5, 31), (1, 3, 3, 1, 5, 49)]          # no masters here: IDs as given


def test_employee_terr_pg_keeps_products_with_no_territory(h):
    r = ok(c.get("/api/terr-pg/employee", headers=h))
    assert r["columns"][:3] == ["Employee Name", "HQ", "Territory"]
    assert r["total"] == 2 and all(x[1] is None and x[2] is None for x in r["rows"])


def test_terr_pg_filters_names_and_downloads(h):
    assert ok(c.get("/api/terr-pg/customer?team_id=1", headers=h))["total"] == 0
    assert ok(c.get("/api/terr-pg/customer?product_group_id=1", headers=h))["total"] == 2
    ok(c.post("/api/masters/customer", headers=h,
              json={"id": 1, "customer_name": "HEALTH MEDECINE DISTRIBUTOR"}), 201)
    r = ok(c.get("/api/terr-pg/customer?q=medecine", headers=h))
    assert r["total"] == 2 and r["rows"][0][0] == "HEALTH MEDECINE DISTRIBUTOR"
    res = c.get("/api/terr-pg/customer/export", headers=h)
    rows = list(openpyxl.load_workbook(io.BytesIO(res.content), read_only=True).active
                .iter_rows(values_only=True))
    assert rows[0][:2] == ("S.No", "Customer Name") and len(rows) == 3
    assert c.get("/api/terr-pg/nobody", headers=h).status_code == 404
