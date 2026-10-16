"""Budget Upload and Budget Report, as the old Abbott screens had them.

    GET  /api/budget/years                  the years of YEAR_MASTER, newest first
    GET  /api/budget/template               AbbottBudgetFormet.xlsx: the header row
    POST /api/budget/upload                 year, file, dry_run: check, then post
    POST /api/budget/carry-forward          0-budget rows: last year's rows missing this year
    POST /api/budget/zero-rows              0-budget rows: year, employee, SKUs, HQs
    POST /api/budget/calculate-actuals      a month's actuals from MAPPING_MASTER(_MID)
    GET  /api/budget/report                 year + optional filters, currency INR | NPR
    GET  /api/budget/report/export          the same, as an Excel workbook

**Upload.** The workbook's first sheet: EMPLOYEE_ID, PRODUCT_GROUP, HQ_NAME,
PRODUCT_BRAND, PRODUCT_ID, PRODUCT_NAME and the twelve <MONTH>_UNITS. Each row
becomes a BUDGET_MASTER row of the year chosen, as the old procedure
(AbbottExecuteSPCreateAbtBudgetUpload) made it: the units, each month's value
as units x the SKU's NRV, PRODUCT_GROUP_ID / PRODUCT_BRAND_ID / HQ_ID, and the
brand's REPORTING as PRODUCT_BRAND_CATEGORY.

* Checked first, and all or nothing: a row with an unknown employee, SKU or
  HQ, a SKU with no NRV, units that are not a number, or the same employee /
  HQ / SKU twice, stops the upload, each named with its row.
* The SKU decides the brand and group (IDs and names, from the masters). A file
  naming another brand or group, or an employee not assigned the HQ, is said
  but does not stop it.
* Safe to send again: a row is the year + employee + HQ + SKU; one already
  there has its units and values replaced, its actuals kept. S_NO is the
  budget line of a SKU in a year, as in the old data: shared by its HQs and
  employees, a new one for a SKU new to the year.
"""
from __future__ import annotations

import io
from collections import Counter
from datetime import date, datetime

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from .. import models_master as X
from ..database import get_db
from ..services import master_io as IO
from ..services.names import norm
from .masters import rate_on

router = APIRouter(prefix="/api")

MONTHS = X.MONTHS
COLUMNS = ["EMPLOYEE_ID", "PRODUCT_GROUP", "HQ_NAME", "PRODUCT_BRAND", "PRODUCT_ID", "PRODUCT_NAME",
           *[f"{m}_UNITS" for m in MONTHS]]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(header: list[str], rows: list[list], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/budget/years")
def years(db: Session = Depends(get_db)):
    return [y for (y,) in db.execute(select(X.YearMaster.year_number)
                                     .order_by(X.YearMaster.year_number.desc()))]


@router.get("/budget/template")
def template():
    return Response(_xlsx(COLUMNS, [], "Sheet1"), media_type=XLSX,
                    headers={"Content-Disposition": 'attachment; filename="AbbottBudgetFormet.xlsx"'})


# ------------------------------------------------------------------- upload
def _show(v):
    """A cell as a person reads it: 1, not 1.0."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return "" if v is None else v


def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0.0
    if isinstance(v, bool):
        raise ValueError
    return float(str(v).replace(",", "").strip()) if isinstance(v, str) else float(v)


@router.post("/budget/upload")
async def upload(year: int = Form(...), dry_run: bool = Form(True), file: UploadFile = File(...),
                 db: Session = Depends(get_db)):
    if not db.execute(select(X.YearMaster.id).where(X.YearMaster.year_number == year)).first():
        raise HTTPException(422, f"{year} is not in Year.")
    name = file.filename or "budget.xlsx"
    if not name.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Only an .xlsx file is read, as the format: AbbottBudgetFormet.xlsx.")
    data = await file.read()
    try:
        sheets = IO.read_tables(name, data)
    except Exception as e:
        raise HTTPException(422, f"The file could not be opened: {e}")
    # the first sheet with the header row
    want = {norm(c): c for c in COLUMNS}
    found = None
    for _, title, rows, _ in sheets:
        for i, r in enumerate(rows[:10]):
            heads = {norm(v): j for j, v in enumerate(r) if v is not None}
            if norm("EMPLOYEE_ID") in heads and norm("PRODUCT_ID") in heads:
                found = (title, i, {want[k]: j for k, j in heads.items() if k in want}, rows)
                break
        if found:
            break
    if not found:
        raise HTTPException(422, "No sheet has the header row of the format (EMPLOYEE_ID … "
                                 "PRODUCT_ID … JANUARY_UNITS …).")
    title, hr, col, rows = found
    missing = [c for c in COLUMNS if c not in col]
    if missing:
        raise HTTPException(422, f"The header row has no {', '.join(missing)}.")

    emps = {i: n for i, n in db.execute(select(X.EmployeeMaster.id, X.EmployeeMaster.employee_ft_nm))}
    P, B, G, H = X.ProductMaster, X.ProductMasterBrand, X.ProductGroup, X.Headquarter
    prods = {r.id: r for r in db.execute(select(P.id, P.sap_product_name, P.nrv, P.product_brand_id,
                                                P.product_group_id))}
    brands = {r.id: r for r in db.execute(select(B.id, B.product_brand_name, B.product_group_id,
                                                 B.reporting))}
    groups = dict(db.execute(select(G.legacy_id, G.name)).all())
    hqs = {norm(n): (lid, n) for lid, n in db.execute(select(H.legacy_id, H.name)) if lid is not None}
    hq_ids = {lid: (lid, n) for lid, n in hqs.values()}
    covers = set(db.execute(select(X.EmployeeTerrAssignment.employee_id, X.EmployeeTerrAssignment.hq_id)))

    errors, warnings, out = [], [], []
    seen: dict[tuple, int] = {}
    warn_count: Counter = Counter()
    for n, r in enumerate(rows[hr + 1:], start=hr + 2):
        cell = lambda c: r[col[c]] if col[c] < len(r) else None     # noqa: E731
        if all(v in (None, "") for v in r):
            continue
        bad = []
        try:
            emp = int(_num(cell("EMPLOYEE_ID")))
        except ValueError:
            emp = None
        if not emp or emp not in emps:
            bad.append(f"EMPLOYEE_ID {_show(cell('EMPLOYEE_ID'))} is not in Employee")
        try:
            pid = int(_num(cell("PRODUCT_ID")))
        except ValueError:
            pid = None
        p = prods.get(pid)
        if p is None:
            bad.append(f"PRODUCT_ID {_show(cell('PRODUCT_ID'))} is not in SKU")
        elif not p.nrv:
            bad.append(f"SKU {pid} ({p.sap_product_name}) has no NRV, so its value cannot be worked out")
        hq_raw = cell("HQ_NAME")
        hq = hqs.get(norm(hq_raw))
        if hq is None and isinstance(hq_raw, (int, float)):
            hq = hq_ids.get(int(hq_raw))
        if hq is None:
            bad.append(f"HQ_NAME {_show(hq_raw)} is not in Head Quarter")
        units = []
        for m in MONTHS:
            try:
                u = _num(cell(f"{m}_UNITS"))
                if u < 0:
                    raise ValueError
                units.append(u)
            except (ValueError, TypeError):
                bad.append(f"{m}_UNITS should be a number of 0 or more, not '{_show(cell(f'{m}_UNITS'))}'")
                units.append(0.0)
        if not bad:
            key = (emp, hq[0], pid)
            if key in seen:
                bad.append(f"the same employee, HQ and SKU as row {seen[key]}")
            seen[key] = n
        if bad:
            errors.append({"row": n, "problems": bad})
            continue
        b = brands.get(p.product_brand_id)
        gid = p.product_group_id if p.product_group_id is not None else (b.product_group_id if b else None)
        gname, bname = groups.get(gid), (b.product_brand_name if b else None)
        if cell("PRODUCT_BRAND") and bname and norm(cell("PRODUCT_BRAND")) != norm(bname):
            warn_count["brand"] += 1
            if warn_count["brand"] <= 5:
                warnings.append(f"Row {n}: the file says brand {cell('PRODUCT_BRAND')}, SKU {pid} is "
                                f"{bname}; the SKU's is used.")
        if cell("PRODUCT_GROUP") and gname and norm(cell("PRODUCT_GROUP")) != norm(gname):
            warn_count["group"] += 1
            if warn_count["group"] <= 5:
                warnings.append(f"Row {n}: the file says group {cell('PRODUCT_GROUP')}, SKU {pid} is "
                                f"in {gname}; the SKU's is used.")
        if covers and (emp, hq[0]) not in covers:
            warn_count["hq"] += 1
            if warn_count["hq"] <= 5:
                warnings.append(f"Row {n}: employee {emp} ({emps[emp]}) has no territory in {hq[1]} "
                                f"(Employee Terr-PG).")
        nrv = float(p.nrv)
        out.append({"employee_id": emp, "hq_id": hq[0], "hq_name": hq[1], "product_id": pid,
                    "product_name": p.sap_product_name, "product_brand_id": p.product_brand_id,
                    "product_brand": bname, "product_group_id": gid, "product_group": gname,
                    "product_brand_category": b.reporting if b else None,
                    **{f"{m.lower()}_units": u for m, u in zip(MONTHS, units)},
                    **{m.lower(): u * nrv for m, u in zip(MONTHS, units)}})
    for k, what in (("brand", "brand"), ("group", "group"), ("hq", "HQ")):
        if warn_count[k] > 5:
            warnings.append(f"… and {warn_count[k] - 5} more rows with another {what}.")

    Bm = X.Budget
    existing = {}
    for r in db.execute(select(Bm.id, Bm.employee_id, Bm.hq_id, Bm.product_id)
                        .where(Bm.target_year == year)):
        existing.setdefault((r.employee_id, r.hq_id, r.product_id), r.id)
    lines = dict(db.execute(select(Bm.product_id, func.min(Bm.s_no))
                            .where(Bm.target_year == year).group_by(Bm.product_id)).all())
    added = sum(1 for o in out if (o["employee_id"], o["hq_id"], o["product_id"]) not in existing)
    result = {"dry_run": dry_run, "year": year, "sheet": title, "rows": len(out) + len(errors),
              "ok_rows": len(out), "added": added, "updated": len(out) - added,
              "errors": errors[:200], "error_rows": len(errors), "warnings": warnings,
              "total_units": round(sum(sum(o[f"{m.lower()}_units"] for m in MONTHS) for o in out), 2),
              "total_value": round(sum(sum(o[m.lower()] for m in MONTHS) for o in out), 2)}
    if errors:
        result["message"] = (f"Nothing was posted: {len(errors)} row(s) to put right first."
                             if not dry_run else f"{len(errors)} row(s) to put right before posting.")
        return result
    if dry_run:
        return result
    next_line = (db.execute(select(func.max(Bm.s_no))).scalar() or 0) + 1
    for o in out:
        rid = existing.get((o["employee_id"], o["hq_id"], o["product_id"]))
        if rid is not None:
            row = db.get(Bm, rid)
        else:
            s_no = lines.get(o["product_id"])
            if s_no is None:
                s_no = lines[o["product_id"]] = next_line
                next_line += 1
            row = Bm(s_no=s_no, target_year=year)
            db.add(row)
        for k, v in o.items():
            setattr(row, k, v)
    db.commit()
    result["message"] = f"Posted: {added} row(s) added, {len(out) - added} updated, for {year}."
    return result


# ------------------------------------------------------------- zero budget rows
def _zero_rows(db: Session, year: int, keys: set, dry: bool) -> dict:
    """Make a 0-budget row for each (employee, HQ Master ID, SKU) not yet in the
    year: units and values 0, the SKU's brand and group, the HQ's name, the
    brand's REPORTING; S_NO the SKU's line in the year, or a new one. A row
    already there is left as it is; a SKU, HQ or employee not in the masters is
    left out and counted."""
    Bm, P, B, G, H = X.Budget, X.ProductMaster, X.ProductMasterBrand, X.ProductGroup, X.Headquarter
    have = set(db.execute(select(Bm.employee_id, Bm.hq_id, Bm.product_id).where(Bm.target_year == year)))
    prods = {r.id: r for r in db.execute(select(P.id, P.sap_product_name, P.product_brand_id, P.product_group_id))}
    brands = {r.id: r for r in db.execute(select(B.id, B.product_brand_name, B.product_group_id, B.reporting))}
    groups = dict(db.execute(select(G.legacy_id, G.name)).all())
    hqs = dict(db.execute(select(H.legacy_id, H.name).where(H.legacy_id.is_not(None))).all())
    emps = set(db.execute(select(X.EmployeeMaster.id)).scalars())
    new, there, unknown = [], 0, Counter()
    for emp, hq, pid in sorted(keys):
        if (emp, hq, pid) in have:
            there += 1
        elif pid not in prods:
            unknown["SKU"] += 1
        elif hq not in hqs:
            unknown["HQ"] += 1
        elif emp not in emps:
            unknown["employee"] += 1
        else:
            new.append((emp, hq, pid))
    if not dry and new:
        lines = dict(db.execute(select(Bm.product_id, func.min(Bm.s_no))
                                .where(Bm.target_year == year).group_by(Bm.product_id)).all())
        next_line = (db.execute(select(func.max(Bm.s_no))).scalar() or 0) + 1
        for emp, hq, pid in new:
            p = prods[pid]
            b = brands.get(p.product_brand_id)
            gid = p.product_group_id if p.product_group_id is not None else (b.product_group_id if b else None)
            if pid not in lines:
                lines[pid] = next_line
                next_line += 1
            db.add(Bm(s_no=lines[pid], target_year=year, employee_id=emp, hq_id=hq, hq_name=hqs[hq],
                      product_id=pid, product_name=p.sap_product_name, product_brand_id=p.product_brand_id,
                      product_brand=b.product_brand_name if b else None, product_group_id=gid,
                      product_group=groups.get(gid), product_brand_category=b.reporting if b else None,
                      **{f"{m.lower()}_units": 0 for m in MONTHS}, **{m.lower(): 0 for m in MONTHS}))
        db.commit()
    return {"dry_run": dry, "year": year, "added": len(new), "already_there": there,
            "left_out": dict(unknown)}


def _year(db: Session, y, what: str) -> int:
    try:
        y = int(y)
    except (TypeError, ValueError):
        raise HTTPException(422, f"Choose the {what}.")
    if not db.execute(select(X.YearMaster.id).where(X.YearMaster.year_number == y)).first():
        raise HTTPException(422, f"{y} is not in Year.")
    return y


@router.post("/budget/carry-forward")
def carry_forward(body: dict = Body(...), db: Session = Depends(get_db)):
    """Every employee / HQ / SKU budgeted in from_year and missing from to_year
    gets a 0-budget row in to_year (or only one employee's, when given)."""
    src, dst = _year(db, body.get("from_year"), "year to copy from"), _year(db, body.get("to_year"), "year")
    if src == dst:
        raise HTTPException(422, "Choose two different years.")
    Bm = X.Budget
    sel = select(Bm.employee_id, Bm.hq_id, Bm.product_id).where(
        Bm.target_year == src, Bm.employee_id.is_not(None), Bm.hq_id.is_not(None), Bm.product_id.is_not(None))
    if body.get("employee_id"):
        sel = sel.where(Bm.employee_id == int(body["employee_id"]))
    keys = set(db.execute(sel))
    return {**_zero_rows(db, dst, keys, bool(body.get("dry_run"))), "from_year": src}


@router.post("/budget/zero-rows")
def zero_rows(body: dict = Body(...), db: Session = Depends(get_db)):
    """A 0-budget row for a new SKU (or several): year, employee, SKUs, and HQs —
    when none are chosen, the HQs the employee covers on Employee Terr-PG."""
    year = _year(db, body.get("year"), "year")
    try:
        emp = int(body.get("employee_id"))
    except (TypeError, ValueError):
        raise HTTPException(422, "Choose the employee.")
    pids = sorted({int(x) for x in body.get("product_ids") or []})
    if not pids:
        raise HTTPException(422, "Choose the product (one or more).")
    hqs = sorted({int(x) for x in body.get("hq_ids") or []})
    if not hqs:
        hqs = sorted(set(db.execute(select(X.EmployeeTerrAssignment.hq_id)
                                    .where(X.EmployeeTerrAssignment.employee_id == emp)).scalars()))
        if not hqs:
            raise HTTPException(422, "Choose the HQ: the employee has none on Employee Terr-PG.")
    keys = {(emp, h, p) for h in hqs for p in pids}
    return {**_zero_rows(db, year, keys, bool(body.get("dry_run"))), "hq_ids": hqs}


# ------------------------------------------------------------ calculate actuals
# The old Calculate Actuals (SP_CALCULATE_ACTUALS_NEW, SP_CALCULATE_RETURN_VALUE_NEW
# and AbbottCalculateActualsWithMultiplyAndUpdateSQL), for one month of a year,
# over every budget row of the year. A row's sales are MAPPING_MASTER's rows of
# the month with its HQ_ID, PRODUCT_GROUP_ID, PRODUCT_BRAND_ID, PRODUCT_ID and
# PRODUCT_BRAND_CATEGORY (not its employee: each employee of the HQ gets them).
#   <M>_ACTUAL                SUM(RATE x (SALES_QTY + RETURN_QUANTITY)) / 100000 (lakh, NPR)
#   <M>_ACTUAL_UNITS          SUM(SALES_QTY + RETURN_QUANTITY)
#   <M>_MID_ACTUAL(_UNITS)    the same from MAPPING_MASTER_MID
#   <M>_RETURN                SUM(RETURN_QUANTITY x RATE) / the INR/NPR rate (INR); empty
#                             when the month has no such rows
# Values (not units) are halved for product group 4 in HQ 6. As in SQL, a row
# whose RATE or RETURN_QUANTITY is empty counts for nothing. Then the year's
# group names METABOLICS / VACCINE become METABOLICS-NEW / VACCINE-NEURO, and an
# empty <M>_ACTUAL is 0. Every row is updated by its own ID, so a row with no
# PRODUCT_BRAND_CATEGORY is never doubled (the old procedure inserted another).
KEY = ("hq_id", "product_group_id", "product_brand_id", "product_id", "product_brand_category")
HALVED = (4, 6)                      # (PRODUCT_GROUP_ID, HQ_ID)
RENAMED = {"METABOLICS": "METABOLICS-NEW", "VACCINE": "VACCINE-NEURO"}


def _sums(db: Session, T, month: str, year: int, returns: bool) -> dict:
    cols = [getattr(T, k) for k in KEY]
    q = T.sales_qty + T.return_quantity
    parts = [func.sum(T.rate * q), func.sum(q), func.count()]
    if returns:
        parts.append(func.sum(T.return_quantity * T.rate))
    out = {}
    for r in db.execute(select(*cols, *parts).where(T.month == month, T.year == year).group_by(*cols)):
        out[tuple(r[:5])] = r[5:]
    return out


@router.post("/budget/calculate-actuals")
def calculate_actuals(body: dict = Body(...), db: Session = Depends(get_db)):
    year = _year(db, body.get("year"), "year")
    month = str(body.get("month") or "").strip().title()
    if month.upper() not in MONTHS:
        raise HTTPException(422, "Choose the month.")
    m = month.lower()
    try:
        inr = float(rate_on(db, "INR", "NPR", date(year, MONTHS.index(month.upper()) + 1, 1))["rate"])
    except HTTPException as e:
        raise HTTPException(422, f"{e.detail} The returns are worked out in INR: add the rate on Currency rates.")
    main = _sums(db, X.SalesMapping, month, year, returns=True)
    mid = _sums(db, X.MidSalesMapping, month, year, returns=False)
    Bm = X.Budget
    rows = db.execute(select(Bm).where(Bm.target_year == year)).scalars().all()
    used, with_sales, empty_key = set(), 0, 0
    tot = {"actual": 0.0, "units": 0.0, "mid_actual": 0.0, "mid_units": 0.0, "returns": 0.0}
    for r in rows:
        key = tuple(getattr(r, k) for k in KEY)
        half = 2.0 if (r.product_group_id, r.hq_id) == HALVED else 1.0
        if any(v is None for v in key):      # as SQL: an empty column matches nothing
            empty_key += 1
            a = b = None
        else:
            a, b = main.get(key), mid.get(key)
        if a:
            used.add(key)
        if b:
            used.add(key)
        actual = (float(a[0]) / 100000 / half) if a and a[0] is not None else 0.0
        units = float(a[1]) if a and a[1] is not None else 0.0
        mid_actual = (float(b[0]) / 100000 / half) if b and b[0] is not None else 0.0
        mid_units = float(b[1]) if b and b[1] is not None else 0.0
        ret = (float(a[3]) / inr / half) if a and a[3] is not None else None
        setattr(r, f"{m}_actual", actual)
        setattr(r, f"{m}_actual_units", units)
        setattr(r, f"{m}_mid_actual", mid_actual)
        setattr(r, f"{m}_mid_actual_units", mid_units)
        setattr(r, f"{m}_return", ret)
        if actual or units or mid_actual or mid_units:
            with_sales += 1
        tot["actual"] += actual
        tot["units"] += units
        tot["mid_actual"] += mid_actual
        tot["mid_units"] += mid_units
        tot["returns"] += ret or 0.0
        if r.product_group in RENAMED:
            r.product_group = RENAMED[r.product_group]
    db.flush()
    # the old procedure's last step, for the rows the loop did not reach
    for old, new in RENAMED.items():
        db.execute(update(Bm).where(Bm.target_year == year, Bm.product_group == old).values(product_group=new))
    db.execute(update(Bm).where(Bm.target_year == year, getattr(Bm, f"{m}_actual").is_(None))
               .values({getattr(Bm, f"{m}_actual"): 0}))
    db.commit()
    unbudgeted = [k for k in set(main) | set(mid) if k not in used and None not in k]
    return {"year": year, "month": month, "rows": len(rows), "rows_with_sales": with_sales,
            "inr_rate": inr, **{k: round(v, 4) for k, v in tot.items()},
            "sales_without_budget_row": len(unbudgeted), "rows_with_an_empty_key": empty_key,
            "message": (f"{month} {year}: actuals worked out on {len(rows)} budget rows, "
                        f"{with_sales} of them with sales."
                        + (f" {len(unbudgeted)} HQ / product sales of the month have no budget row: "
                           f"give them one on Zero Budget and calculate again." if unbudgeted else "")
                        + (f" {empty_key} budget rows have an empty HQ, group, brand, product or "
                           f"category ID and get no actuals." if empty_key else ""))}


# ------------------------------------------------------------------- report
# As the old Budget Report (BudgetReport_<date>.xls): these columns, one row per
# budget row whose SKU is in the SKU master (the old report joined it), then a
# Total row. BUDGET_MASTER's values are NPR; the report gives them in INR at the
# INR/NPR Currency rate in force today (or NPR, as stored, when asked).
FILTERS = ("employee_id", "hq_id", "product_group_id", "product_brand_id", "product_id")
REPORT_COLUMNS = ["TARGET_YEAR", "EMPLOYEE_ID", "EMPLOYEE_NAME", "HQ_ID", "HQ_NAME", "PRODUCT_GROUP_ID",
                  "PRODUCT_GROUP", "PRODUCT_BRAND_ID", "PRODUCT_BRAND", "PRODUCT_ID", "PRODUCT_NAME",
                  *[f"{m}_UNITS" for m in MONTHS], *MONTHS]


def _report(db: Session, params: dict) -> dict:
    try:
        year = int(params.get("year"))
    except (TypeError, ValueError):
        raise HTTPException(422, "Choose a year.")
    currency = (params.get("currency") or "INR").upper()
    if currency not in ("INR", "NPR"):
        raise HTTPException(422, "Currency is INR or NPR.")
    rate, note = 1.0, "Values in NPR, as BUDGET_MASTER keeps them."
    if currency == "INR":
        r = rate_on(db, "INR", "NPR", date.today())
        rate = float(r["rate"])
        note = (f"Values in INR: BUDGET_MASTER's NPR divided by {rate:g}, the INR/NPR rate in force "
                f"from {r['effective_from']} (Currency rates).")
    Bm, P = X.Budget, X.ProductMaster
    sel = select(Bm).join(P, P.id == Bm.product_id).where(Bm.target_year == year)
    for f in FILTERS:
        v = params.get(f)
        if v not in (None, ""):
            try:
                sel = sel.where(getattr(Bm, f) == int(v))
            except ValueError:
                raise HTTPException(422, f"{f} should be a number.")
    rows = db.execute(sel.order_by(Bm.employee_id, Bm.hq_name, Bm.product_group, Bm.product_brand,
                                   Bm.product_name)).scalars().all()
    # rows naming a SKU the master does not have: the old report left them out too
    left = db.execute(select(func.count()).select_from(Bm).outerjoin(P, P.id == Bm.product_id)
                      .where(Bm.target_year == year, P.id.is_(None))).scalar()
    ids = {r.employee_id for r in rows}
    names = {e.id: e.full_name for e in db.execute(
        select(X.EmployeeMaster).where(X.EmployeeMaster.id.in_(ids))).scalars()} if ids else {}
    out = []
    tu, tv = [0.0] * 12, [0.0] * 12
    for r in rows:
        units = [float(getattr(r, f"{m.lower()}_units") or 0) for m in MONTHS]
        values = [float(getattr(r, m.lower()) or 0) / rate for m in MONTHS]
        for k in range(12):
            tu[k] += units[k]
            tv[k] += values[k]
        out.append({"id": r.id, "target_year": r.target_year, "employee_id": r.employee_id,
                    "employee_name": names.get(r.employee_id), "hq_id": r.hq_id, "hq_name": r.hq_name,
                    "product_group_id": r.product_group_id, "product_group": r.product_group,
                    "product_brand_id": r.product_brand_id, "product_brand": r.product_brand,
                    "product_id": r.product_id, "product_name": r.product_name,
                    "units": units, "values": values})
    return {"year": year, "currency": currency, "rate": rate, "note": note, "left_out": left,
            "months": [m.title() for m in MONTHS], "rows": out,
            "total_units": tu, "total_values": tv}


@router.get("/budget/report")
def report(request: Request, db: Session = Depends(get_db)):
    r = _report(db, dict(request.query_params))
    r["rows"] = [{**x, "units": [round(v, 4) for v in x["units"]],
                  "values": [round(v, 4) for v in x["values"]]} for x in r["rows"]]
    return r


@router.get("/budget/report/export")
def report_export(request: Request, db: Session = Depends(get_db)):
    """The old report's workbook: its columns, its Total row."""
    r = _report(db, dict(request.query_params))
    body = [[x["target_year"], x["employee_id"], x["employee_name"], x["hq_id"], x["hq_name"],
             x["product_group_id"], x["product_group"], x["product_brand_id"], x["product_brand"],
             x["product_id"], x["product_name"], *x["units"], *x["values"]] for x in r["rows"]]
    body.append(["Total", *[None] * 10, *r["total_units"], *r["total_values"]])
    stamp = datetime.now().strftime("%m_%d_%Y_%H_%M")
    return Response(_xlsx(REPORT_COLUMNS, body, f"Budget {r['year']} {r['currency']}"), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="BudgetReport_{stamp}.xlsx"'})
