"""Customer Terr-PG and Employee Terr-PG: the old system's assignment screens.

The old application showed one page per kind, where each row is a product
assignment with its territory beside it:

    Customer Name | HQ | Territory | Product Group | Team | Product Brand | Product
    Employee Name | HQ | Territory | Product Group | Product Brand | Product

The data is kept as it was in the old system, in two tables each:

* customer: CUSTOMER_PG_ASSIGNMENT (customer, group, brand, product) joined to
  CUSTOMER_TERR_ASSIGNMENT (customer, territory, HQ, group, team) on the
  customer *and* the product group — a customer is covered by a different
  territory and team for each group.
* employee: EMPLOYEE_PG_ASSIGNMENT (employee, product, brand, group) joined to
  EMPLOYEE_TERR_ASSIGNMENT (employee, HQ, territory) on the employee.

A product row with no territory row still shows, with HQ and Territory empty,
so nothing assigned is hidden. IDs are shown by name; an ID no master knows is
shown as the number.

    GET /api/terr-pg/{kind}            kind = customer | employee; q, page, size, filters
    GET /api/terr-pg/{kind}/export     the same rows as an Excel workbook
"""
from __future__ import annotations

import io
from dataclasses import dataclass

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from .. import models_master as X
from ..database import get_db
from ..masters import specs as S
from ..masters.view import name_of

router = APIRouter(prefix="/api")


@dataclass
class Col:
    label: str
    side: str          # "pg" or "terr": which table the ID is on
    attr: str
    ref: str           # the master it names (a spec slug)
    key: str = "id"    # that master's column the ID matches


KINDS = {
    "customer": dict(
        title="Customer Terr-PG", who="customer_id", who_ref="customer",
        pg=X.CustomerPgAssignment, terr=X.CustomerTerrAssignment,
        on=("customer_id", "product_group_id"),
        cols=[Col("Customer Name", "pg", "customer_id", "customer"),
              Col("HQ", "terr", "hq_id", "headquarter", "legacy_id"),
              Col("Territory", "terr", "territory_id", "territory", "legacy_id"),
              Col("Product Group", "pg", "product_group_id", "product_group", "legacy_id"),
              Col("Team", "terr", "team_id", "team", "legacy_id"),
              Col("Product Brand", "pg", "product_brand_id", "brand"),
              Col("Product", "pg", "product_id", "sku")],
        filters=[("customer_id", "pg"), ("hq_id", "terr"), ("product_group_id", "pg"),
                 ("team_id", "terr")]),
    "employee": dict(
        title="Employee Terr-PG", who="employee_id", who_ref="employee",
        pg=X.EmployeePgAssignment, terr=X.EmployeeTerrAssignment,
        on=("employee_id",),
        cols=[Col("Employee Name", "pg", "employee_id", "employee"),
              Col("HQ", "terr", "hq_id", "headquarter", "legacy_id"),
              Col("Territory", "terr", "territory_id", "territory", "legacy_id"),
              Col("Product Group", "pg", "product_group_id", "product_group", "legacy_id"),
              Col("Product Brand", "pg", "product_brand_id", "brand"),
              Col("Product", "pg", "product_id", "sku")],
        filters=[("employee_id", "pg"), ("hq_id", "terr"), ("product_group_id", "pg")]),
}


def _kind(kind: str) -> dict:
    k = KINDS.get(kind)
    if k is None:
        raise HTTPException(404, "No such assignment page")
    return k


def _select(k: dict, params: dict, q: str | None):
    pg, terr = k["pg"], k["terr"]
    cols = [getattr(pg if c.side == "pg" else terr, c.attr) for c in k["cols"]]
    on = and_(*[getattr(pg, a) == getattr(terr, a) for a in k["on"]])
    sel = select(*cols).select_from(pg).outerjoin(terr, on)
    for name, side in k["filters"]:
        v = params.get(name)
        if v not in (None, ""):
            try:
                sel = sel.where(getattr(pg if side == "pg" else terr, name) == int(v))
            except ValueError:
                raise HTTPException(422, f"{name} should be a number.")
    if q and q.strip():
        like = f"%{q.strip()}%"
        who = S.get(k["who_ref"]).model
        name_col = who.customer_name if k["who_ref"] == "customer" \
            else who.employee_ft_nm + " " + func.coalesce(who.employee_lt_nm, "")
        sel = sel.where(or_(
            getattr(pg, k["who"]).in_(select(who.id).where(name_col.like(like))),
            pg.product_id.in_(select(X.ProductMaster.id)
                              .where(X.ProductMaster.sap_product_name.like(like))),
            pg.product_brand_id.in_(select(X.ProductMasterBrand.id)
                                    .where(X.ProductMasterBrand.product_brand_name.like(like)))))
    order = [getattr(pg, k["who"]), pg.product_group_id, pg.product_brand_id, pg.product_id]
    return sel, order


def _names(db: Session, k: dict, rows) -> list[list]:
    """Every ID of the page by name, one query per column."""
    out = [list(r) for r in rows]
    for ci, c in enumerate(k["cols"]):
        target = S.get(c.ref)
        col = getattr(target.model, c.key)
        ids = {r[ci] for r in out} - {None}
        found = {}
        if ids:
            for t in db.execute(select(target.model).where(col.in_(ids))).scalars():
                found[getattr(t, c.key)] = name_of(target, t)
        for r in out:
            v = r[ci]
            r[ci] = found.get(v, v) if v is not None else None
    return out


@router.get("/terr-pg/{kind}")
def list_rows(kind: str, request: Request, q: str | None = None, page: int = 1, size: int = 50,
              db: Session = Depends(get_db)):
    k = _kind(kind)
    size = max(1, min(size, 500))
    sel, order = _select(k, dict(request.query_params), q)
    total = db.execute(select(func.count()).select_from(sel.subquery())).scalar()
    rows = db.execute(sel.order_by(*order).offset((max(page, 1) - 1) * size).limit(size)).all()
    return {"title": k["title"], "columns": [c.label for c in k["cols"]],
            "filters": [{"name": n, "ref": next(c.ref for c in k["cols"] if c.attr == n),
                         "key": next(c.key for c in k["cols"] if c.attr == n),
                         "label": next(c.label for c in k["cols"] if c.attr == n)}
                        for n, _ in k["filters"]],
            "total": total, "page": page, "size": size, "rows": _names(db, k, rows)}


@router.get("/terr-pg/{kind}/export")
def export(kind: str, request: Request, q: str | None = None, db: Session = Depends(get_db)):
    k = _kind(kind)
    sel, order = _select(k, dict(request.query_params), q)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(k["title"].replace("/", "-"))
    head = []
    for c in ["S.No", *[c.label for c in k["cols"]]]:
        cell = WriteOnlyCell(ws, value=c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        head.append(cell)
    ws.append(head)
    rows = db.execute(sel.order_by(*order)).all()
    n = 0
    for i in range(0, len(rows), 2000):                  # names a chunk at a time
        for r in _names(db, k, rows[i:i + 2000]):
            n += 1
            ws.append([n, *r])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = k["title"].replace(" ", "_").replace("/", "-").upper() + ".xlsx"
    return StreamingResponse(buf, headers={"Content-Disposition": f'attachment; filename="{fname}"'},
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
