"""Reading the distributors' workbooks, and the PDF reports some of them send.

Two layouts turn up. A cross-tab - products down the page, customers across,
quantities on one row and amounts on the next - and a plain row-wise sheet with
a header naming its columns. Which one a file uses is worked out from the file,
not from its name, because the names are not reliable.

A PDF is read the same way once its tables are lifted out: every table on every
page is joined into one list of rows, a header repeated at the top of each page
is dropped, and the rows go through the same two layouts. Report PDFs often
print the item (or the party) once as a group heading and leave it blank on the
lines below, so in a PDF a blank product or customer takes the one above it.
"""
from __future__ import annotations

import io
import re
import zipfile

from openpyxl import load_workbook

# "40 + 0" is 40 bought and 0 free. A minus may be written apart from the
# figure ("- 40 + 0"); it stays with it, so that is -40 bought.
QTY = re.compile(r"^\s*(-?\s*[\d.,]+)\s*\+\s*(-?\s*[\d.,]+)\s*$")
BOOK = re.compile(r"\.(xlsx|xlsm|xls)$", re.I)
PDF = re.compile(r"\.pdf$", re.I)

WANT = {
    "customer": ["customername", "customer", "party", "partyname", "shopname"],
    "product": ["productname", "product", "item", "itemname", "particulars", "prodname",
                "itemdescription", "productdescription"],
    "qty": ["quantity", "qty"],
    "free": ["freequantity", "freeqty", "free", "fqty"],
    "amount": ["amount", "netamount", "value", "bamount", "basicamount"],
    "invoice": ["invoiceno", "invoicenumber", "invno", "billno", "invoice"],
    "rate": ["rate", "salerate", "salesrate", "unitrate", "ptr"],
}


def clean(s) -> str:
    """Drop special characters. Punctuation becomes a space rather than nothing,
    so 'Rantoks Pharma Pvt.ltd.' reads 'Rantoks Pharma Pvt ltd', not 'Pvtltd'."""
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9 ]+", " ", str(s if s is not None else ""))).strip()


def _norm(s) -> str:
    return re.sub(r"[^a-z]", "", str(s if s is not None else "").lower())


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[\s,]", "", str(v if v is not None else "")))
    except ValueError:
        return 0.0


def readable(name: str) -> bool:
    return bool(BOOK.search(name) or PDF.search(name))


def _skip_entry(inner: str) -> bool:
    """Zip entries that are not files anybody sent: folders and the resource
    forks a Mac adds (__MACOSX/, ._name, .DS_Store)."""
    base = inner.split("/")[-1]
    return (inner.endswith("/") or inner.startswith("__MACOSX/") or not base
            or base.startswith("._") or base == ".DS_Store")


def _unique(base: str, taken: set) -> str:
    """Two folders in one zip can hold files with the same name. Both are
    kept: the second becomes 'name (2).pdf' rather than overwriting the first."""
    if base not in taken:
        taken.add(base)
        return base
    stem, dot, ext = base.rpartition(".")
    if not dot:
        stem, ext = base, ""
    i = 2
    while True:
        cand = f"{stem} ({i}).{ext}" if ext else f"{stem} ({i})"
        if cand not in taken:
            taken.add(cand)
            return cand
        i += 1


def books_in(name: str, data: bytes, everything: bool = False, taken: set | None = None):
    """Whatever was uploaded, reduced to the files this can read: a zip is
    opened, and the workbooks and PDFs inside it - or sent loose - are taken
    as they are. Anything else is left out.

    With everything=True (distributors whose files are chosen by hand under all
    conditions, e.g. Gunjeshwari) every file extracted from the zip is offered,
    whatever its type or name, and a zip inside the zip is opened too. Whether
    a file can be read is then decided when it is read, not here.

    `taken` is shared across everything uploaded together, so two zips sent at
    once (Gunjeshwari's ALI and AHL zips) that hold files with the same name
    keep both: the second becomes 'name (2).pdf' rather than overwriting."""
    out = []
    taken = set() if taken is None else taken

    def take(inner_name: str, blob: bytes, depth: int):
        base = inner_name.split("/")[-1]
        if everything and depth < 2 and base.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(blob)) as z:
                    for info in sorted(z.infolist(), key=lambda i: i.filename):
                        if info.is_dir() or _skip_entry(info.filename):
                            continue
                        take(info.filename, z.read(info), depth + 1)
                return
            except zipfile.BadZipFile:
                pass                      # offered as it is; reading it will say why
        if everything:
            out.append((_unique(base, taken), blob))
            return
        if not readable(base) or base.startswith(".") or base.startswith("_"):
            return
        out.append((_unique(base, taken), blob))

    if name.lower().endswith(".zip"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for info in sorted(z.infolist(), key=lambda i: i.filename):
                    if info.is_dir() or _skip_entry(info.filename):
                        continue
                    take(info.filename, z.read(info), 1)
        except zipfile.BadZipFile:
            return [(_unique(name.split("/")[-1], taken), data)] if everything else []
        return out
    if everything or readable(name):
        return [(_unique(name.split("/")[-1], taken), data)]
    return []


def _cross_tab_header(rows) -> int:
    for r in range(min(len(rows), 25)):
        cells = [_norm(c) for c in (rows[r] or [])[:6]]
        if "prodname" in cells and "sn" in cells and "unit" in cells:
            return r
    return -1


def _parse_cross_tab(rows, hr, out, stats):
    head = rows[hr]
    cols = [c for c in range(5, len(head))
            if head[c] is not None and str(head[c]).strip() and _norm(head[c]) != "total"]
    r = hr + 1
    while r < len(rows):
        row = rows[r]
        sn = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        prod = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        if sn != "1" or not prod or _norm(prod) == "total":
            r += 1
            continue
        amounts = rows[r + 1] if r + 1 < len(rows) else []
        for c in cols:
            cell = row[c] if c < len(row) else None
            m = QTY.match(str(cell if cell is not None else ""))
            if not m:
                continue
            q, f = _num(m.group(1)), _num(m.group(2))
            if q == 0 and f == 0:            # nothing bought and nothing given
                stats["omitted"] += 1
                continue
            amt = _num(amounts[c]) if c < len(amounts) else 0.0
            out.append({"customer": clean(head[c]), "product": clean(prod),
                        "qty": q, "free": f, "amount": amt, "invoice": None, "rate": None,
                        "is_return": False})
        r += 2                               # the amount row has been consumed
    return "cross-tab"


def _map_columns(rows):
    for r in range(min(len(rows), 30)):
        cells = [_norm(c) for c in (rows[r] or [])]
        found = {}
        for key, names in WANT.items():
            for i, c in enumerate(cells):
                if c and c in names:
                    found[key] = i
                    break
        if all(k in found for k in ("customer", "product", "qty", "amount")):
            return r, found
    return None


def _parse_row_wise(rows, hr, cmap, out, stats, is_return, negate, fill_down=False):
    sign = -1 if (is_return and negate) else 1
    last = {"customer": "", "product": ""}
    for r in range(hr + 1, len(rows)):
        row = rows[r]

        def get(k):
            return row[cmap[k]] if k in cmap and cmap[k] < len(row) else None

        cust = str(get("customer") or "").strip()
        prod = str(get("product") or "").strip()
        if fill_down:
            # a group heading row (item or party alone, no figures) only sets
            # what the lines below it belong to
            if (cust or prod) and get("qty") in (None, "") and get("amount") in (None, ""):
                if cust:
                    last["customer"] = cust
                if prod:
                    last["product"] = prod
                continue
            if not cust and not prod:
                # A line with figures but no names is either the same product
                # carried on to another batch (it starts from the quantity, and
                # carries a rate) or a subtotal (no rate). Only the first is a sale.
                if get("rate") in (None, "") or not (last["customer"] or last["product"]):
                    continue
            if _norm(cust).endswith("total") or _norm(prod).endswith("total"):
                continue
            cust = cust or last["customer"]
            prod = prod or last["product"]
            last["customer"], last["product"] = cust, prod
        if not cust and not prod:
            continue
        if _norm(cust) in ("total", "grandtotal") or _norm(prod) in ("total", "grandtotal"):
            continue
        m = QTY.match(str(get("qty") or ""))
        if m:
            q, f = _num(m.group(1)), _num(m.group(2))
        else:
            q, f = _num(get("qty")), _num(get("free"))
        if q == 0 and f == 0:
            stats["omitted"] += 1
            continue
        inv = get("invoice")
        rate = get("rate")
        out.append({"customer": clean(cust), "product": clean(prod),
                    "qty": q * sign, "free": f * sign, "amount": _num(get("amount")) * sign,
                    "invoice": str(inv).strip() if inv not in (None, "") else None,
                    # the rate is a price, so a return keeps it positive
                    "rate": _num(rate) if rate not in (None, "") else None,
                    "is_return": bool(is_return)})
    return "rows (sales return)" if is_return else "rows"


def _pdf_rows(data: bytes):
    """Every table row on every page, as text. Ruled tables are read from their
    lines; a report printed without lines is read from the gaps between words."""
    import pdfplumber

    def cell(v):
        return re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for settings in (None, {"vertical_strategy": "text", "horizontal_strategy": "text"}):
            rows = []
            for page in pdf.pages:
                tables = page.extract_tables(settings) if settings else page.extract_tables()
                for t in tables:
                    rows.extend([[cell(c) for c in row] for row in t if row])
            if _cross_tab_header(rows) >= 0 or _map_columns(rows):
                return rows, len(pdf.pages)
        return [], len(pdf.pages)


def read_pdf(name: str, data: bytes, negate_returns: bool = True) -> dict:
    """Every sales row in a PDF report, read as one sheet."""
    out, stats = [], {"omitted": 0}
    try:
        rows, pages = _pdf_rows(data)
    except Exception as e:
        return {"name": name, "rows": [], "omitted": 0, "sheets": [],
                "error": f"Not a PDF this can open: {type(e).__name__}"}
    how = None
    hr = _cross_tab_header(rows)
    if hr >= 0:
        how = _parse_cross_tab(rows, hr, out, stats)
    else:
        found = _map_columns(rows)
        if found:
            hr, cmap = found
            head = [_norm(c) for c in rows[hr]]
            # the header printed again at the top of every page is not a sales row
            body = [r for r in rows[hr + 1:] if [_norm(c) for c in r] != head]
            how = _parse_row_wise([rows[hr]] + body, 0, cmap, out, stats,
                                  bool(re.search(r"return", name, re.I)), negate_returns,
                                  fill_down=True)
    sheets = [{"name": f"PDF, {pages} page{'s' if pages != 1 else ''}",
               "layout": how and f"{how}, from PDF", "rows": len(out),
               "omitted": stats["omitted"]}]
    error = None if out else (
        "No sales rows could be read from this PDF. It needs a table with a header row naming "
        "customer (or party), product (or item), quantity and amount. A scanned PDF has no "
        "text to read.")
    return {"name": name, "rows": out, "omitted": stats["omitted"], "sheets": sheets,
            "error": error}


def _sheets(data: bytes) -> list[tuple[str, list[list]]]:
    """Every sheet as (title, rows of values). An .xlsx is opened with openpyxl;
    an older .xls (the binary format Yetichem's AHCPW.xls, AILGPW.xls and
    AILNPW.xls come in) with xlrd. Which one is decided by trying, because a
    file named .xls is sometimes an .xlsx underneath."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        import xlrd                          # .xls (BIFF); raises if it is not one either
        bk = xlrd.open_workbook(file_contents=data)

        def val(c):
            if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                return None
            v = c.value
            # xlrd gives every number as a float; 1.0 must read as 1, as in .xlsx
            if c.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                return int(v)
            return v

        return [(sh.name, [[val(c) for c in sh.row(r)] for r in range(sh.nrows)])
                for sh in bk.sheets()]
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    finally:
        wb.close()


# ====================================================================
# Reading a file somebody chose by hand
#
# For a distributor whose files are chosen by hand under all conditions
# (Gunjeshwari), a file is read whatever its name or extension, and if the
# strict reading above finds nothing, two looser readings are tried in turn:
#
#   1. a table whose header names the columns in other words (Party / A/c
#      Name, Item / Description, Qty / Sale Qty, Sch / Bonus, Amt / Value, PTR
#      / Price ...). Customer and Amount may be missing: the party is then taken
#      from a heading line above the items, and Amount from Quantity x Rate;
#   2. for a PDF, the plain text, line by line: a line ending in figures is a
#      sale, a line of words above it names the party.
#
# The rows then go through the same distributor rules as any other file.
# ====================================================================

LOOSE = {
    "customer": ("customername", "customer", "partyname", "party", "acname", "accountname",
                 "account", "ledgername", "ledger", "retailer", "retailername", "buyer",
                 "buyername", "billedto", "billto", "shopname", "shop", "chemist",
                 "chemistname", "doctorname", "client", "clientname", "dealer", "dealername",
                 "stockist", "outlet", "name"),
    "product": ("productname", "product", "itemname", "item", "itemdescription",
                "productdescription", "description", "particulars", "particular", "prodname",
                "medicine", "medicinename", "brand", "brandname", "skuname", "sku",
                "material", "materialname", "goods", "itemdesc", "proddesc"),
    "qty": ("quantity", "qty", "saleqty", "salesqty", "sqty", "billqty", "billedqty", "soldqty",
            "salequantity", "salesquantity", "netqty", "nos", "units", "pcs", "qtysold",
            "quantitysold", "sales"),
    "free": ("freequantity", "freeqty", "free", "fqty", "sch", "schqty", "scheme", "schemeqty",
             "bonus", "bonusqty", "freegoods", "fq", "schfree"),
    "rate": ("rate", "salerate", "salesrate", "unitrate", "ptr", "pts", "price", "unitprice",
             "netrate", "billrate", "srate"),
    "amount": ("amount", "amt", "netamount", "netamt", "value", "netvalue", "bamount",
               "basicamount", "salevalue", "salesvalue", "saleamount", "salesamount",
               "totalamount", "totalvalue", "grossamount", "grossvalue", "taxable",
               "taxablevalue", "taxableamount", "billamount", "total"),
}
# "name" and "total" alone are too common to trust unless nothing better is on
# the header row, so they are tried last.
_WEAK = {"name", "total", "sales", "nos", "units"}
_NOT_A_NAME = re.compile(r"(^|\b)(grand\s*total|sub\s*total|total|page\s*\d|continued)", re.I)
_LABEL = re.compile(r"^\s*(party|customer|a/?c|account|ledger|retailer|buyer|chemist|"
                    r"dealer|shop|client)(\s*name)?\s*[:\-.]\s*", re.I)
_NUMTOK = re.compile(r"^\(?-?\s*(rs\.?|inr|\u20b9)?\s*-?[\d,]*\.?\d+\)?(cr|dr)?$", re.I)


def _isnum(v) -> bool:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return True
    t = str(v if v is not None else "").strip()
    return bool(t) and bool(_NUMTOK.match(t.replace(" ", "")))


def _tonum(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"(?i)rs\.?|inr|\u20b9|cr$|dr$|[\s,]", "", str(v or ""))
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        x = float(t)
    except ValueError:
        return 0.0
    return -x if neg else x


def _key_of(cell) -> str | None:
    n = _norm(cell)
    if not n:
        return None
    for k, names in LOOSE.items():
        if n in names and n not in _WEAK:
            return k
    for k, names in LOOSE.items():           # 'Sale Qty (Strips)' -> qty
        for w in names:
            if len(w) >= 3 and w not in _WEAK and n.startswith(w):
                return k
    for k, names in LOOSE.items():
        if n in names:
            return k
    return None


def _loose_header(rows):
    """The first row in the top 40 that names at least a product (or a party)
    column and a quantity column, in any of the usual words."""
    for r in range(min(len(rows), 40)):
        cmap = {}
        for i, c in enumerate(rows[r] or []):
            k = _key_of(c)
            if k and k not in cmap:
                cmap[k] = i
        if "qty" in cmap and ("product" in cmap or "customer" in cmap):
            if "product" not in cmap:         # only a name column: it is the item
                cmap["product"] = cmap.pop("customer")
            return r, cmap
    return None


def _heading_text(row) -> str:
    parts = [str(c).strip() for c in row if c not in (None, "") and str(c).strip()]
    return _LABEL.sub("", " ".join(parts)).strip()


def _parse_loose(rows, out, stats, is_return, negate) -> str | None:
    found = _loose_header(rows)
    if not found:
        return None
    hr, cmap = found
    head = [_norm(c) for c in rows[hr]]
    sign = -1 if (is_return and negate) else 1
    last = {"customer": "", "product": ""}
    before = len(out)
    for row in rows[hr + 1:]:
        row = list(row or [])
        if [_norm(c) for c in row] == head:  # header repeated on the next page
            continue

        def get(k):
            return row[cmap[k]] if k in cmap and cmap[k] < len(row) else None

        cust = str(get("customer") or "").strip()
        prod = str(get("product") or "").strip()
        qcell = get("qty")
        m = QTY.match(str(qcell if qcell is not None else ""))
        has_qty = bool(m) or _isnum(qcell)
        has_amt = _isnum(get("amount"))
        if not has_qty and not has_amt:
            # a heading: it sets the party (or item) the lines below belong to
            text = _heading_text(row)
            if not text or _NOT_A_NAME.search(text) or not re.search(r"[A-Za-z]", text):
                continue
            if "customer" in cmap:
                if cust:
                    last["customer"] = _LABEL.sub("", cust).strip()
                if prod:
                    last["product"] = prod
                if not cust and not prod:
                    last["customer"] = text
            else:
                last["customer"] = text   # no party column: a heading names the party
            continue
        if not has_qty:
            continue
        if _NOT_A_NAME.search(cust) or _NOT_A_NAME.search(prod):
            continue                      # Total / Sub Total lines
        if not cust and not prod and "rate" in cmap and not _isnum(get("rate")):
            continue                      # a nameless line with no rate is a subtotal
        if not cust and not prod and "rate" not in cmap:
            continue
        cust = cust or last["customer"]
        prod = prod or last["product"]
        if not prod and not cust:
            continue
        last["customer"], last["product"] = cust, prod
        if m:
            q, f = _num(m.group(1)), _num(m.group(2))
        else:
            q, f = _tonum(qcell), (_tonum(get("free")) if _isnum(get("free")) else 0.0)
        if q == 0 and f == 0:
            stats["omitted"] += 1
            continue
        rate = _tonum(get("rate")) if _isnum(get("rate")) else None
        amt = _tonum(get("amount")) if has_amt else 0.0
        if not amt and rate is not None:
            amt = round(q * rate, 2)      # no amount column: Quantity x Rate
        out.append({"customer": clean(cust), "product": clean(prod),
                    "qty": q * sign, "free": f * sign, "amount": amt * sign,
                    "invoice": None, "rate": rate, "is_return": bool(is_return)})
    return "rows, read loosely" if len(out) > before else None


def _split_line(line: str):
    """'ORS SACHET OR9 12/27 20 0 8.25 165.00' -> ('ORS SACHET', ['20','0','8.25','165.00'])."""
    toks = line.split()
    tail = []
    while toks and _isnum(toks[-1]):
        tail.insert(0, toks.pop())
    # a batch number or an expiry date sits between the name and the figures
    while toks and len(toks) > 1 and (re.fullmatch(r"\d{1,2}[/-]\d{2,4}", toks[-1])
                                      or (re.search(r"\d", toks[-1])
                                          and re.search(r"[A-Za-z]", toks[-1])
                                          and len(toks[-1]) <= 12 and tail)):
        toks.pop()
    # a line that is only a batch number and figures carries on the product above
    if (tail and len(toks) == 1 and re.search(r"\d", toks[0]) and re.search(r"[A-Za-z]", toks[0])
            and len(toks[0]) <= 12):
        toks = []
    return " ".join(toks).strip(), tail


def _text_header(lines):
    """The figure columns named on a header line, in order: qty, free, rate, amount
    and anything else (mrp, disc, gst ...) as None so the positions still line up."""
    for ln in lines[:60]:
        words = re.findall(r"[A-Za-z.%/]+", ln)
        keys = [_key_of(w) for w in words]
        if "qty" in keys and ("amount" in keys or "rate" in keys):
            cols = []
            for w, k in zip(words, keys):
                n = _norm(w)
                if k in ("qty", "free", "rate", "amount"):
                    cols.append(k)
                elif n in ("mrp", "disc", "discount", "gst", "tax", "cgst", "sgst", "igst",
                           "exp", "expiry", "pack", "packing"):
                    cols.append(None)
            return ln, cols
    return None, None


def _assign(tail, cols):
    nums = [_tonum(t) for t in tail]
    if cols and len(cols) == len(nums):
        got = {k: v for k, v in zip(cols, nums) if k}
        if "qty" in got:
            return got
    n = len(nums)
    if n == 1:
        return {"qty": nums[0]}
    if n == 2:
        return {"qty": nums[0], "amount": nums[1]}
    if n == 3:
        return {"qty": nums[0], "rate": nums[1], "amount": nums[2]}
    # four or more: Qty, Free, Rate ... Amount (last)
    return {"qty": nums[0], "free": nums[1], "rate": nums[-2], "amount": nums[-1]}


def _pdf_text(data: bytes, out, stats, is_return, negate) -> str | None:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        lines = []
        for page in pdf.pages:
            lines.extend((page.extract_text() or "").splitlines())
    head, cols = _text_header(lines)
    sign = -1 if (is_return and negate) else 1
    party, product, before = "", "", len(out)
    for ln in lines:
        ln = ln.strip()
        if not ln or ln == head:
            continue
        name, tail = _split_line(ln)
        if _NOT_A_NAME.search(name):
            continue
        if not tail:
            # words only: a party heading (or a report title, which is harmless
            # because the first real party heading replaces it)
            text = _LABEL.sub("", ln).strip()
            if re.search(r"[A-Za-z]", text) and len(text) <= 120:
                party, product = text, ""
            continue
        if not name and not product:
            continue
        got = _assign(tail, cols)
        if not name and "rate" not in got:
            continue                       # a nameless line with no rate is a subtotal
        product = name or product
        q, f = got.get("qty", 0.0), got.get("free", 0.0)
        if q == 0 and f == 0:
            stats["omitted"] += 1
            continue
        rate = got.get("rate")
        amt = got.get("amount") or (round(q * rate, 2) if rate is not None else 0.0)
        out.append({"customer": clean(party), "product": clean(product),
                    "qty": q * sign, "free": f * sign, "amount": amt * sign,
                    "invoice": None, "rate": rate, "is_return": bool(is_return)})
    return "text lines" if len(out) > before else None


def _csv_rows(data: bytes):
    import csv
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.reader(io.StringIO(text), dialect)]


def _kind(name: str, data: bytes) -> str:
    """What a file is, from its first bytes rather than its name."""
    head = data[:8]
    if head.startswith(b"%PDF") or b"%PDF" in data[:1024]:
        return "pdf"
    if head.startswith(b"PK"):
        return "book"                       # .xlsx / .xlsm
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "book"                       # .xls
    if PDF.search(name):
        return "pdf"
    if BOOK.search(name):
        return "book"
    if re.search(r"\.(csv|tsv|txt|prn)$", name, re.I):
        return "text"
    try:
        data[:2048].decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        return "other"


def _read_any(name: str, data: bytes, negate: bool) -> dict:
    """A hand-picked file, read under all conditions."""
    kind = _kind(name, data)
    is_ret = bool(re.search(r"return", name, re.I))
    if kind == "pdf":
        r = read_pdf(name, data, negate)
        if r["rows"]:
            return r
        out, stats = [], {"omitted": 0}
        how, pages = None, 0
        try:
            rows, pages = _pdf_rows_all(data)
            how = _parse_loose(rows, out, stats, is_ret, negate)
            if not how:
                out, stats = [], {"omitted": 0}
                how = _pdf_text(data, out, stats, is_ret, negate)
        except Exception as e:
            return {**r, "error": f"Not a PDF this can open: {type(e).__name__}"}
        sheets = [{"name": f"PDF, {pages} page{'s' if pages != 1 else ''}",
                   "layout": how and f"{how}, from PDF", "rows": len(out),
                   "omitted": stats["omitted"]}]
        return {"name": name, "rows": out, "omitted": stats["omitted"], "sheets": sheets,
                "error": None if out else (
                    "No sales lines could be found in this PDF, even read loosely. If it is a "
                    "scanned image there is no text in it to read.")}
    if kind == "book":
        r = read_book(name if BOOK.search(name) else name + ".xlsx", data, negate)
        r["name"] = name
        if r["rows"]:
            return r
        try:
            book = _sheets(data)
        except Exception:
            return r
        out, stats, sheets = [], {"omitted": 0}, []
        for title, rows in book:
            before, ob = len(out), stats["omitted"]
            how = _parse_loose(rows, out, stats,
                               is_ret or bool(re.search(r"return", title, re.I)), negate)
            sheets.append({"name": title, "layout": how, "rows": len(out) - before,
                           "omitted": stats["omitted"] - ob})
        return {"name": name, "rows": out, "omitted": stats["omitted"], "sheets": sheets,
                "error": None if out else r["error"]}
    if kind == "text":
        try:
            rows = _csv_rows(data)
        except Exception as e:
            return {"name": name, "rows": [], "omitted": 0, "sheets": [],
                    "error": f"Not a text file this can open: {type(e).__name__}"}
        out, stats = [], {"omitted": 0}
        hr = _cross_tab_header(rows)
        if hr >= 0:
            how = _parse_cross_tab(rows, hr, out, stats)
        else:
            found = _map_columns(rows)
            how = (_parse_row_wise(rows, found[0], found[1], out, stats, is_ret, negate,
                                   fill_down=True) if found else None)
            if not out:
                out, stats = [], {"omitted": 0}
                how = _parse_loose(rows, out, stats, is_ret, negate)
        return {"name": name, "rows": out, "omitted": stats["omitted"],
                "sheets": [{"name": "text", "layout": how, "rows": len(out),
                            "omitted": stats["omitted"]}],
                "error": None if out else (
                    "No sales rows could be read from this file. It needs a header row "
                    "naming at least the product (or item) and the quantity.")}
    return {"name": name, "rows": [], "omitted": 0, "sheets": [],
            "error": "This is not a PDF, Excel or text file, so there is nothing in it to read."}


def _pdf_rows_all(data: bytes):
    """Every table row on every page, from both the ruled and the text table
    strategies, whichever names its columns loosely first."""
    import pdfplumber

    def cell(v):
        return re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for settings in (None, {"vertical_strategy": "text", "horizontal_strategy": "text"}):
            rows = []
            for page in pdf.pages:
                tables = page.extract_tables(settings) if settings else page.extract_tables()
                for t in tables:
                    rows.extend([[cell(c) for c in row] for row in t if row])
            if _loose_header(rows):
                return rows, len(pdf.pages)
        return [], len(pdf.pages)


def read_book(name: str, data: bytes, negate_returns: bool = True,
              any_file: bool = False) -> dict:
    """Every sales row in one workbook (or PDF report), with a note of how each
    sheet was read.

    any_file=True is for a file chosen by hand under all conditions: it is
    read whatever its name or type, and more loosely when the strict reading
    finds nothing (see _read_any)."""
    if any_file:
        return _read_any(name, data, negate_returns)
    if PDF.search(name):
        return read_pdf(name, data, negate_returns)
    out, stats, sheets = [], {"omitted": 0}, []
    try:
        book = _sheets(data)
    except Exception as e:
        return {"name": name, "rows": [], "omitted": 0, "sheets": [],
                "error": f"Not a workbook this can open: {type(e).__name__}"}
    for title, rows in book:
        before, omitted_before = len(out), stats["omitted"]
        hr = _cross_tab_header(rows)
        if hr >= 0:
            how = _parse_cross_tab(rows, hr, out, stats)
        else:
            found = _map_columns(rows)
            how = (_parse_row_wise(rows, found[0], found[1], out, stats,
                                   bool(re.search(r"return", title, re.I)),
                                   negate_returns) if found else None)
        sheets.append({"name": title, "layout": how,
                       "rows": len(out) - before, "omitted": stats["omitted"] - omitted_before})
    error = None if out else (
        "No sales rows could be read. Expected either the cross-tab layout (Sn / Prodname / "
        "Unit) or a header row naming customer, product, quantity and amount.")
    return {"name": name, "rows": out, "omitted": stats["omitted"], "sheets": sheets,
            "error": error}
