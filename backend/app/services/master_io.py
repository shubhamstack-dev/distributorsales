"""Masters in and out of Excel: the reading half of "upload master data" and
the whole of the exports.

Every master is described once in masters/specs.py, so one reader serves
them all:

* **Which master a sheet is** comes from its name — the table name in
  MASTER DATA & TABLES.xlsx (CUSTOMER_MASTER, PRODUCT_MASTER_BRAND …), the old
  system's (ZONE, HQ_MASTER, PRODUCT_MASTER_GROUP …), or the screen's title.
  A workbook with one sheet is named by its file.
* **Which column is which** comes from the header: the workbook's column name,
  the field's label on the screen, or one of the old names the spec lists.
* **A reference** is written the way people know it: a Master ID (a number),
  a code, or a name. The IDs in the Excel tables are Master IDs.

The export writes the same headers back, so an export is also the template.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

NULLS = {"", "NULL", "NONE", "NAN", "N/A", "#N/A"}
# columns read past on purpose, with the reason given in the report
IGNORED = {"PASSWORD": "PASSWORD is never stored, so it was read past."}


def norm(s) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(s if s is not None else "").upper())


def is_excel_table(spec) -> bool:
    """The four tables kept exactly as MASTER DATA & TABLES.xlsx has them."""
    return spec.model.__tablename__.isupper()


# ------------------------------------------------------------- the columns
def columns(spec) -> list[tuple[str, str]]:
    """(header, attribute) in the order an export writes them."""
    if is_excel_table(spec):
        mapper = spec.model.__mapper__
        out = []
        for col in spec.model.__table__.columns:
            attr = next(p.key for p in mapper.column_attrs if p.columns[0] is col)
            out.append((col.name, attr))
        return out
    out = []
    for f in spec.fields:
        if f.derived:
            continue
        out.append((f.label, f.name))
    return out


def header_map(spec) -> dict[str, str]:
    """norm(header) -> attribute, for every name a column may go by."""
    m: dict[str, str] = {}
    mapper = spec.model.__mapper__
    # weakest first, so a more specific name written later wins
    for f in spec.fields:
        if f.type == "aliases":
            continue
        m[norm(f.name)] = f.name
        m[norm(f.label)] = f.name
    allowed = {f.name for f in spec.fields if not f.derived and f.type != "aliases"}
    alias = next((f for f in spec.fields if f.type == "aliases"), None)
    if alias:
        allowed |= {f"alias_{i}" for i in range(1, alias.count + 1)}
    for p in mapper.column_attrs:
        if p.key in allowed:
            m[norm(p.columns[0].name)] = p.key
    for f in spec.fields:
        for c in f.cols or []:
            m[norm(c)] = f.name
    return m


def sheet_keys(spec) -> set[str]:
    return {norm(x) for x in (spec.slug, spec.title, spec.model.__tablename__, *spec.sheets)}


def detect(specs, *names) -> object | None:
    for n in names:
        k = norm(n)
        if not k:
            continue
        for s in specs:
            if k in sheet_keys(s):
                return s
    return None


# ---------------------------------------------------------------- reading
def _cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return None if v.upper() in NULLS else v
    return v


def read_tables(name: str, data: bytes) -> list[tuple[str, str, list[list], int]]:
    """(file stem, sheet title, rows, sheets in the file) for every sheet."""
    stem = re.sub(r"\.[^.]+$", "", re.split(r"[\\/]", name)[-1])
    low = name.lower()
    if low.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            out = []
            for ws in wb.worksheets:
                rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
                out.append((stem, ws.title, rows, len(wb.worksheets)))
            return out
        finally:
            wb.close()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    first = (text or "").split("\n", 1)[0]
    delim = "\t" if first.count("\t") > first.count(",") else ","
    rows = [[_cell(v) for v in r] for r in csv.reader(io.StringIO(text or ""), delimiter=delim)]
    return [(stem, "", rows, 1)]


def find_header(rows, hmap) -> tuple[int, dict[int, str], list[str]]:
    """The first of the top rows naming at least two of the master's columns.
    Returns (row index, {column index: attribute}, headers not used)."""
    for r in range(min(len(rows), 6)):
        got, unused = {}, []
        for i, h in enumerate(rows[r] or []):
            if h is None:
                continue
            k = norm(h)
            if k in hmap and hmap[k] not in got.values():
                got[i] = hmap[k]
            else:
                unused.append(str(h))
        if len(got) >= 2:
            return r, got, unused
    return -1, {}, []


# ------------------------------------------------------------- the values
def to_int(v):
    if v is None or isinstance(v, bool):
        return None if v is None else int(v)
    if isinstance(v, (int, float, Decimal)):
        if float(v) != int(float(v)):
            raise ValueError
        return int(v)
    s = str(v).strip().replace(",", "")
    return int(float(s))


def to_text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        v = v.date()
    s = " ".join(str(v).split())
    return s or None


def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:19]
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y", "%d/%m/%Y",
                "%m/%d/%Y", "%d-%b-%Y", "%d-%b-%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s if (" " in fmt or "T" in fmt) else s[:11].strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError


def to_bool(v, default=1):
    if v is None:
        return default
    if isinstance(v, bool):
        return int(v)
    s = str(v).strip().lower()
    if s in ("1", "1.0", "true", "yes", "y", "active", "open", "current"):
        return 1
    if s in ("0", "0.0", "false", "no", "n", "inactive", "closed"):
        return 0
    raise ValueError


def to_decimal(v):
    if v is None:
        return None
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError


# --------------------------------------------------- natural keys of a row
def _name_cols(spec):
    m = spec.model
    for a in ("name", "customer_name", "product_brand_name", "sap_product_name", "label"):
        if hasattr(m, a) and a in m.__mapper__.attrs:
            return getattr(m, a)
    return None


def natural_key(spec, o):
    """What a row is called when another master's export points at it."""
    if o is None:
        return None
    if spec.master_id:
        return o.legacy_id
    if is_excel_table(spec):
        return o.id
    if spec.slug == "year":
        return o.year_value
    if hasattr(o, "code") and getattr(o, "code", None):
        return o.code
    if hasattr(o, "name"):
        return o.name
    return o.id


def find_row(db: Session, spec, value, year_id=None):
    """The row a value names: a Master ID or the row's own ID (numbers), else a
    code or a name, case ignored."""
    if value is None:
        return None
    m = spec.model
    num = None
    try:
        num = to_int(value)
    except (TypeError, ValueError):
        pass
    if num is not None:
        if spec.master_id:
            r = db.execute(select(m).where(m.legacy_id == num)).scalars().first()
            if r is not None:
                return r
        elif is_excel_table(spec):
            return db.get(m, num)
        elif spec.slug == "year":
            r = db.execute(select(m).where(m.year_value == num)).scalars().first()
            if r is not None:
                return r
    text = to_text(value).upper()
    if "code" in m.__mapper__.attrs:
        q = select(m).where(func.upper(m.code) == text)
        if year_id is not None and "year_id" in m.__mapper__.attrs:
            q = q.where(m.year_id == year_id)
        r = db.execute(q).scalars().first()
        if r is not None:
            return r
    col = _name_cols(spec)
    if col is not None:
        r = db.execute(select(m).where(func.upper(col) == text)).scalars().first()
        if r is not None:
            return r
    if spec.slug == "employee":
        full = func.upper(func.trim(m.employee_ft_nm + " " + func.coalesce(m.employee_lt_nm, "")))
        return db.execute(select(m).where(full == text)).scalars().first()
    return None


def natural_id(db: Session, spec, value):
    r = find_row(db, spec, value)
    return r.id if r is not None else None


# ---------------------------------------------------------------- export
HEAD_FILL = PatternFill("solid", fgColor="0B1626")


def _write_sheet(ws, db: Session, spec, specs_by_slug):
    cols = columns(spec)
    for i, (h, _) in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center")
    fields = {f.name: f for f in spec.fields}
    # refs written by their natural key, looked up once per target
    lookups = {}
    for f in spec.fields:
        if f.type == "ref" and not f.key and not f.soft:
            t = specs_by_slug[f.ref]
            lookups[f.name] = {o.id: natural_key(t, o)
                               for o in db.execute(select(t.model)).scalars()}
    rows = db.execute(select(spec.model).order_by(spec.model.id)).scalars().all()
    for r_i, o in enumerate(rows, 2):
        for c_i, (_, attr) in enumerate(cols, 1):
            v = getattr(o, attr)
            if attr in lookups and v is not None:
                v = lookups[attr].get(v, v)
            elif isinstance(v, Decimal):
                v = float(v)
            f = fields.get(attr)
            cell = ws.cell(r_i, c_i, v)
            if isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
            if f is not None and f.type == "bool" and v is not None:
                cell.value = int(v)
    for i, (h, _) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(36, len(str(h)) + 4))
    ws.freeze_panes = "A2"
    return len(rows)


def _sheet_title(spec) -> str:
    t = spec.model.__tablename__ if is_excel_table(spec) else spec.title
    return re.sub(r"[:\\/?*\[\]]", "-", t)[:31]


def export(db: Session, specs, only=None) -> bytes:
    by_slug = {s.slug: s for s in specs}
    wb = Workbook()
    wb.remove(wb.active)
    for s in specs:
        if only and s.slug not in only:
            continue
        _write_sheet(wb.create_sheet(_sheet_title(s)), db, s, by_slug)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
