"""Writing the sheet the distributors' rules describe."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADS = ["Distributor Name", "Invoice No", "Month", "Year", "Mid Month", "Customer Name",
         "Product Name", "Quantity", "Free Quantity", "Rate", "B.Amount", "Amount"]
ARIAL = "Arial"


def build(batch, rows, files) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    thin = Side(style="thin", color="D3D4E4")
    fill = PatternFill("solid", fgColor="2B3990")
    for i, h in enumerate(HEADS, 1):
        c = ws.cell(1, i, h)
        c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for n, r in enumerate(rows, start=1):
        i = n + 1
        vals = [r.distributor_name, r.invoice_no, r.month_name, int(r.year), r.mid_month,
                r.customer_name, r.product_name, float(r.quantity), float(r.free_quantity),
                float(r.rate), float(r.b_amount), float(r.amount)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            cell.font = Font(name=ARIAL, size=10)
            cell.border = Border(bottom=thin)
            if c == 4:
                cell.number_format = "0"
                cell.alignment = Alignment(horizontal="center")
            elif c == 5:
                cell.alignment = Alignment(horizontal="center")
            elif c in (8, 9):
                cell.number_format = "#,##0"
            elif c in (10, 11, 12):
                cell.number_format = "#,##0.00"

    for col, w in zip("ABCDEFGHIJKL", [18, 14, 12, 8, 11, 34, 28, 10, 14, 12, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{len(rows) + 1}"

    d = batch.distributor
    mid = batch.mid_cutoff_day or batch.cutoff_day
    sheets = bool(d.sales_sheet or d.return_sheet)
    used = [f for f in files if f.used]
    left = [f for f in files if not f.used]
    notes = [
        ("How this sheet was built", ""),
        ("", ""),
        ("Distributor", batch.distributor.name),
        ("Uploaded", batch.source_name),
        ("Files converted", ", ".join(f.file_name for f in used) or "none"),
        ("Files offered but not converted", ", ".join(f.file_name for f in left) or "none"),
        ("Produced for the date", batch.as_of.strftime("%d %B %Y")),
        ("Month cut-off day", str(batch.cutoff_day)),
        ("Month", f"{batch.month_name} - day {batch.as_of.day} is "
                  f"{'on or before' if batch.as_of.day <= batch.cutoff_day else 'after'} the "
                  f"{batch.cutoff_day}th"),
        ("Year", str(batch.year)),
        ("Mid Month", f"{batch.mid_month} - day {batch.as_of.day} is "
                      f"{'on or before' if batch.mid_month == 'N' else 'after'} the {mid}th"),
        ("Invoice No", "Read from the file where present, otherwise numbered in sequence."
                       if batch.invoice_mode == "file"
                       else "INV - 1 upwards, one running number across every file."),
        ("Rate", ("Read from the file. Where a line has no rate, Amount divided by Quantity."
                  if (batch.distributor.rate_mode or "calc") == "file" else
                  "Amount divided by Quantity. Zero where quantity is zero, so a free-only line "
                  "cannot divide by zero.")),
        ("B.Amount and Amount", ("Each read from its own column of the file (where the file has "
                                 "only one, both carry it). Sales positive, sales returns "
                                 "negative." if sheets else "Both read from the file and identical.")),
        ("Sheets read", (f"Sheet {d.sales_sheet} of each workbook as sales"
                         + (f", sheet {d.return_sheet} as sales returns (quantity, free, B.Amount "
                            f"and Amount negative)" if d.return_sheet else "") + ".")
         if sheets else "Every sheet; a sheet named as a return is negated."),
        ("Names", ("Special characters removed; punctuation became a space. Customer names lose "
                   "every special character; product names keep the '.'."
                   if d.product_keep_dot else
                   "Special characters removed; punctuation became a space, so 'Pvt.ltd.' reads "
                   "'Pvt ltd' rather than 'Pvtltd'.")),
        ("Name mapping and notes", batch.notes or "No reference name lists were applied."),
        ("", ""),
        ("Rows written", str(batch.row_count)),
        ("Rows omitted", f"{batch.omitted_count} - quantity and free quantity both zero"),
        ("Total quantity", str(float(batch.total_qty))),
        ("Total amount", f"{float(batch.total_amount):,.2f}"),
        ("Excluded", "Total columns and Total rows, which are subtotals rather than data."),
        ("Batch", f"#{batch.id}, created {batch.created_at_utc:%Y-%m-%d %H:%M} UTC "
                  f"by {batch.created_by}"),
    ]
    ns = wb.create_sheet("Notes")
    for i, (a, b) in enumerate(notes, 1):
        ca, cb = ns.cell(i, 1, a), ns.cell(i, 2, b)
        ca.font = Font(name=ARIAL, size=11 if i == 1 else 10, bold=(i == 1 or (a and not b)))
        cb.font = Font(name=ARIAL, size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
    ns.column_dimensions["A"].width = 30
    ns.column_dimensions["B"].width = 88

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
