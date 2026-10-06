"""Reference names: replacing extracted customer and product names with the
spelling on a distributor's own list.

Matching ignores case, spaces and every special character, so 'Shree Medical
Hall (P) Ltd.' on the PDF matches 'SHREE MEDICAL HALL P LTD' on the list. A
name with no match is kept as extracted (already cleaned of special
characters) and reported, so the list can be completed rather than the miss
going unnoticed.
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

KINDS = {"customer": ["customername", "customer", "party", "partyname", "name", "billingname"],
         "product": ["productname", "product", "item", "itemname", "skuname", "sku", "name",
                     "particulars"]}


def norm(s) -> str:
    """Letters and digits only, upper case: the key two names are matched on."""
    return re.sub(r"[^A-Z0-9]", "", str(s if s is not None else "").upper())


def _h(s) -> str:
    return re.sub(r"[^a-z]", "", str(s if s is not None else "").lower())


def read_names(data: bytes, kind: str) -> list[str]:
    """Every name in the workbook's name column. The column is the one whose
    header names the kind (Customer Name, Product, Item...); without such a
    header, the first column that holds text."""
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        names: list[str] = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            col, start = None, 0
            for r in range(min(len(rows), 20)):
                heads = [_h(c) for c in rows[r]]
                for want in KINDS[kind]:            # most specific header first
                    if want in heads:
                        col, start = heads.index(want), r + 1
                        break
                if col is not None:
                    break
            if col is None:
                for c in range(max((len(r) for r in rows), default=0)):
                    if any(isinstance(r[c], str) and r[c].strip() for r in rows if c < len(r)):
                        col = c
                        break
            if col is None:
                continue
            for r in rows[start:]:
                v = r[col] if col < len(r) else None
                if isinstance(v, (int, float)):
                    v = str(v)
                if isinstance(v, str) and v.strip() and norm(v):
                    names.append(re.sub(r"\s+", " ", v).strip())
        return names
    finally:
        wb.close()


class Mapper:
    """Looks names up in a reference list and counts what it could not match."""

    def __init__(self, refs: dict[str, str]):
        self.refs = refs                  # norm -> reference spelling
        self.missed: dict[str, int] = {}
        self.hits = 0

    def __call__(self, name: str) -> str:
        if not self.refs:
            return name
        got = self.refs.get(norm(name))
        if got is None:
            self.missed[name] = self.missed.get(name, 0) + 1
            return name
        self.hits += 1
        return got

    def summary(self, kind: str) -> str | None:
        if not self.refs:
            return None
        s = f"{kind.capitalize()} names: {self.hits} rows matched the reference list"
        if self.missed:
            shown = ", ".join(sorted(self.missed)[:15])
            more = len(self.missed) - 15
            s += (f"; {len(self.missed)} name{'s' if len(self.missed) != 1 else ''} not on it and kept "
                  f"as extracted: {shown}{f' and {more} more' if more > 0 else ''}")
        return s + "."
